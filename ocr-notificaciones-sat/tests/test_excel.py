from pathlib import Path
from datetime import datetime, date
import sys, tempfile, unittest, zipfile, json, ast, io, contextlib
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from src import ocr_sat as c
from openpyxl import Workbook, load_workbook
from lxml import etree as ET
from PIL import Image, ImageDraw

def libro_ficticio(path):
    wb=Workbook(); ws=wb.active; ws.title='2026'
    ws.append(['EXPDIENTE','numValorPred','fechEmiPred','fechNotiPred','fecha emision REC','Ffecha de notificacion REC','numValorArbitrio','fechEmiArbi','fechNotiArbi','CodPropietario','CodSigeti','CodSimi','nombre'])
    for number,em,notification in [(1,datetime(2026,9,2),None),(2,'03/09/2026',None),(3,'02/09/2026','10/09/2026'),(4,'02/09/2026','=DATE(2026,9,10)'),(5,'02/09/2026',None),(5,'02/09/2026',None)]:
        ws.append([f'16-012-{number:09d}',None,None,None,em,notification,None,None,None,None,'123',None,'PERSONA FICTICIA'])
    wb.create_sheet('2025')['A1']='SIN CAMBIOS'
    wb.save(path)

class ExcelScopeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.source=Path(self.temp.name)/'original.xlsx'
        self.output=Path(self.temp.name)/'resultado.xlsx'
        libro_ficticio(self.source)
        self.index,_,self.epoch=c.build_index(self.source,('2026',))
    def tearDown(self):self.temp.cleanup()
    def test_valid_match_and_iso_date(self):
        d=dict(expediente='16-012-1-2026',codigo='00123',resolucion='UNO',fecha_notificacion='15/09/2026')
        target,issues=c.evaluar(d,self.index)
        self.assertEqual(issues,[]);self.assertEqual(target['celda'],'F2')
        self.assertEqual(c.fecha(target['emision']),date(2026,9,2))
    def test_scope_and_existing_values_rejected(self):
        for key in [('2025','F2'),('2026','E2'),('2026','F3'),('2026','F4'),('2026','F5'),('2026','F6')]:
            with self.subTest(key=key),self.assertRaises(ValueError):
                c.patch_xlsx(self.source,self.output,{key:date(2026,9,15)},self.epoch)
        with self.assertRaises(ValueError):
            c.patch_xlsx(self.source,self.source,{('2026','F2'):date(2026,9,15)},self.epoch)
    def test_preserves_other_cells_and_zip_parts(self):
        c.patch_xlsx(self.source,self.output,{('2026','F2'):date(2026,9,15)},self.epoch)
        with zipfile.ZipFile(self.source) as a,zipfile.ZipFile(self.output) as b:
            self.assertEqual(a.namelist(),b.namelist())
            changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
            self.assertEqual(set(changed),{'xl/styles.xml','xl/worksheets/sheet1.xml'})
            roots=[ET.fromstring(z.read('xl/worksheets/sheet1.xml')) for z in (a,b)]
            for root in roots:
                cell=root.find('.//s:c[@r="F2"]',c.N)
                if cell is not None:cell.getparent().remove(cell)
            self.assertEqual(ET.tostring(roots[0],method='c14n'),ET.tostring(roots[1],method='c14n'))
        wb=load_workbook(self.output)
        self.assertEqual(wb['2026']['F2'].value.date(),date(2026,9,15))
        self.assertEqual(wb['2026']['F5'].value,'=DATE(2026,9,10)')
        wb.close()

class ConfigurarTests(unittest.TestCase):
    def setUp(self):
        self.original=(c.HOJA,c.FECHA_EMISION)
        self.temp=tempfile.TemporaryDirectory()
        self.source=Path(self.temp.name)/'original.xlsx'
        self.output=Path(self.temp.name)/'resultado.xlsx'
        libro_ficticio(self.source)
    def tearDown(self):
        c.HOJA,c.FECHA_EMISION=self.original
        self.temp.cleanup()
    def test_rejects_bad_values(self):
        for hoja,fecha in [('26',None),(None,'texto'),('2025','02/09/2026'),('abcd',None)]:
            with self.subTest(hoja=hoja,fecha=fecha),self.assertRaises(ValueError):
                c.configurar(hoja,fecha)
        self.assertEqual((c.HOJA,c.FECHA_EMISION),self.original)
    def test_scope_follows_configuration(self):
        c.configurar('2026','03/09/2026')
        index,_,epoch=c.build_index(self.source,('2026',))
        d=dict(expediente='16-012-1-2026',codigo='123',resolucion='UNO',fecha_notificacion='15/09/2026')
        target,issues=c.evaluar(d,index)
        self.assertTrue(any('Fuera del alcance' in i for i in issues))
        d['expediente']='16-012-2-2026'
        target,issues=c.evaluar(d,index)
        self.assertEqual(issues,[]);self.assertEqual(target['celda'],'F3')
        with self.assertRaises(ValueError):
            c.patch_xlsx(self.source,self.output,{('2026','F2'):date(2026,9,15)},epoch)
        c.patch_xlsx(self.source,self.output,{('2026','F3'):date(2026,9,15)},epoch)
        self.assertTrue(self.output.exists())
        with self.assertRaises(ValueError):
            c.build_index(self.source,('2025',))

class MotorFalso:
    nombre='motor ficticio'
    def __init__(self):self.vistas=0
    def leer(self,src):
        self.vistas+=1
        with Image.open(src) as im: im.load()
        return dict(expediente='16-012-1-2026',codigo='123',resolucion='UNO',nombre='PERSONA FICTICIA',fecha_notificacion='15/09/2026',
                    fecha_emision='02/09/2026',fecha_texto=None,observaciones=''),{'motor':self.nombre,'tipo':'REC','roi_fecha':[0.5,0.7,0.7,0.75],
                    'orientacion':'vertical','metodo_recuadro':'rotulos_fecha_y_hora','recuadros':[{'roi':[0.5,0.7,0.7,0.75],'metodo':'rotulos_fecha_y_hora'}]}

class ProyectoBase(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        base=Path(self.temp.name)
        self.fotos=base/'fotos'; (self.fotos/'A'/'sub').mkdir(parents=True); (self.fotos/'B').mkdir()
        for i,rel in enumerate(['A/sub/f1.png','A/f2.png','B/f3.png']):
            Image.new('RGB',(40,30),(i*40,120,200)).save(self.fotos/rel)
        (self.fotos/'B/copia_f3.png').write_bytes((self.fotos/'B/f3.png').read_bytes())
        self.excel=base/'original.xlsx'; libro_ficticio(self.excel)
        self.salida=base/'salida'
    def tearDown(self):self.temp.cleanup()
    def proyecto(self):
        with contextlib.redirect_stdout(io.StringIO()):
            return c.Proyecto(self.fotos,self.excel,self.salida)

class ProyectoTests(ProyectoBase):
    def test_inventory_cache_and_resume(self):
        p=self.proyecto()
        llamadas=[]
        real=c.sha256
        def contador(path):
            if Path(path).suffix=='.png': llamadas.append(path)
            return real(path)
        c.sha256=contador
        try:
            with contextlib.redirect_stdout(io.StringIO()) as out:
                p.inventario()
            self.assertEqual(len(llamadas),4)
            self.assertEqual(len(p.active),3)
            self.assertIn('registros escritos: 3',out.getvalue())
            self.assertTrue((self.salida/'inventario_cache.json').exists())
            llamadas.clear()
            q=self.proyecto()
            with contextlib.redirect_stdout(io.StringIO()) as out:
                q.inventario()
            self.assertEqual(llamadas,[])
            self.assertIn('reutilizadas de la caché: 4',out.getvalue())
            self.assertIn('registros escritos: 0',out.getvalue())
            with contextlib.redirect_stdout(io.StringIO()):
                q.inventario(rehash=True)
            self.assertEqual(len(llamadas),4)
        finally:
            c.sha256=real
        cfg=json.loads((self.salida/'proyecto.json').read_text())
        self.assertEqual(cfg['fecha_emision'],'2026-09-02')
        # Carpetas creadas por la versión anterior (sin fecha_emision) siguen siendo válidas.
        cfg.pop('fecha_emision'); cfg['version']='2026-09-02-v1'
        (self.salida/'proyecto.json').write_text(json.dumps(cfg))
        self.proyecto()
        cfg['excel_sha256']='0'*64
        (self.salida/'proyecto.json').write_text(json.dumps(cfg))
        with self.assertRaises(ValueError):
            self.proyecto()
    def test_sample_processing_and_pending(self):
        p=self.proyecto()
        with contextlib.redirect_stdout(io.StringIO()):
            p.inventario()
        muestra=c.seleccionar_muestra(p.listar_fotos(),p.fotos,cantidad=2)
        ids=p.ids_de(muestra)
        self.assertEqual(len(ids),2)
        motor=MotorFalso()
        with contextlib.redirect_stdout(io.StringIO()):
            resumen=p.procesar(motor,limite=10,solo=ids)
        self.assertEqual(resumen['procesadas'],2); self.assertEqual(motor.vistas,2)
        self.assertEqual(sorted(p.records[k]['estado'] for k in ids),['POR_REVISAR','POR_REVISAR'])
        self.assertEqual(len(p.pendientes()),1)
        with contextlib.redirect_stdout(io.StringIO()):
            resumen=p.procesar(motor,limite=10)
        self.assertEqual(resumen['procesadas'],1); self.assertEqual(p.pendientes(),[])
        self.assertIn('segundos',p.records[ids[0]]['evidencia_ocr'])
        # reprocesar: vuelve a leer las ya leídas conservando la lectura anterior; las aprobadas no.
        p.aprobar(ids[0],dict(expediente='16-012-1-2026',codigo='123',resolucion='UNO',nombre='PERSONA FICTICIA',fecha_notificacion='15/09/2026',fecha_emision='02/09/2026'),'Ana')
        with contextlib.redirect_stdout(io.StringIO()):
            resumen=p.procesar(motor,limite=10,solo=ids,reprocesar=True)
        self.assertEqual(resumen['procesadas'],1); self.assertEqual(p.records[ids[0]]['estado'],'APROBADO')
        self.assertEqual(len(p.records[ids[1]]['historial_ocr']),1); self.assertEqual(p.records[ids[1]]['estado'],'POR_REVISAR')
        with self.assertRaises(ValueError):
            p.ids_de([self.fotos/'no_existe.png'])

class FiabilidadTests(unittest.TestCase):
    def test_vista_tinta_aisla_azul(self):
        im=Image.new('RGB',(200,60),'white'); dr=ImageDraw.Draw(im)
        dr.line((0,55,200,55),fill='black',width=3)            # línea impresa
        dr.text((5,5),'FECHA DE NOTIFICACION',fill='black')    # rótulo impreso
        dr.rectangle((120,15,160,40),fill=(40,60,200))         # trazo de bolígrafo azul
        tinta,fraccion=c.vista_tinta(im)
        self.assertIsNotNone(tinta); self.assertGreater(fraccion,0.05)
        px=tinta.load()
        self.assertEqual(px[140,28],(0,0,0))       # tinta conservada
        self.assertEqual(px[100,55],(255,255,255)) # línea negra eliminada
        self.assertEqual(px[10,8],(255,255,255))   # texto impreso eliminado
        sin,fr=c.vista_tinta(Image.new('RGB',(50,50),'white'))
        self.assertIsNone(sin); self.assertEqual(fr,0.0)
    def test_dates_in_text(self):
        self.assertEqual(c.dates_in_text('15 /09 /2026'),['2026-09-15'])
        self.assertEqual(c.dates_in_text('15 09  2026'),['2026-09-15'])
        self.assertEqual(c.dates_in_text('15092026'),['2026-09-15'])
        self.assertEqual(c.dates_in_text('16/09/26'),['2026-09-16'])
        self.assertEqual(c.dates_in_text('16 09 26'),['2026-09-16'])
        self.assertEqual(c.dates_in_text('21/9/26 y 22/9/26'),['2026-09-21','2026-09-22'])
        self.assertEqual(c.dates_in_text('160926'),[])
        # Espacios entre dígitos: se eliminan sin cambiar cifras (exacta).
        self.assertEqual(c.dates_in_text('1 5 / 0 9 / 2 0 2 6'),['2026-09-15'])
        self.assertEqual(c.extraer_fechas('1 5 / 0 9 / 2 0 2 6'),{'exactas':['2026-09-15'],'corregidas':[]})
        self.assertEqual(c.dates_in_text('5 15/09/2026'),['2026-09-15'])  # el folio no se pega a la fecha
    def test_fechas_corregidas_son_solo_candidatas(self):
        # O→0 dentro de un patrón de fecha: candidata corregida, nunca exacta.
        ex=c.extraer_fechas('15/09/2O26')
        self.assertEqual(ex['exactas'],[]); self.assertEqual([x['fecha'] for x in ex['corregidas']],['2026-09-15'])
        self.assertEqual(ex['corregidas'][0]['regla'],'O→0, I/l→1')
        ex=c.extraer_fechas('l5/O9/2026')
        self.assertEqual(ex['exactas'],[]); self.assertEqual([x['fecha'] for x in ex['corregidas']],['2026-09-15'])
        # Año de 5 cifras: candidata ambigua.
        ex=c.extraer_fechas('15/09/12026')
        self.assertEqual(ex['exactas'],[]); self.assertEqual([x['fecha'] for x in ex['corregidas']],['2026-09-15'])
        self.assertEqual(ex['corregidas'][0]['regla'],'año de 5 cifras')
        # Fuera de un patrón de fecha no se sustituye nada; no se completan dígitos ni se adivina el año.
        self.assertEqual(c.extraer_fechas('FOLIOS 5 OI lO'),{'exactas':[],'corregidas':[]})
        self.assertEqual(c.extraer_fechas('15/09/202'),{'exactas':[],'corregidas':[]})
        self.assertEqual(c.extraer_fechas('15/09'),{'exactas':[],'corregidas':[]})
        self.assertEqual(c.extraer_fechas('OO/OO/OOOO'),{'exactas':[],'corregidas':[]})
        # Una exacta y su versión corregida son la misma fecha: no se duplica.
        ex=c.extraer_fechas('15/09/2026 15/09/2O26')
        self.assertEqual(ex,{'exactas':['2026-09-15'],'corregidas':[]})
        # Las corregidas no cuentan para la confianza: solo entran las 'fechas' exactas.
        v=[{'vista':'trocr_color','fechas':[],'fechas_corregidas':ex['corregidas']},{'vista':'vlm_color','fechas':[]}]
        self.assertEqual(c.nivel_confianza(v,'2026-09-15'),('MEDIA',[]))
        self.assertEqual(c.dates_in_text('Siendo las 9:00 horas, del dia 11/09/2026'),['2026-09-11'])
        self.assertEqual(c.dates_in_text('15-9-2026 y 03/06/2026'),['2026-06-03','2026-09-15'])
        self.assertEqual(c.dates_in_text('31022026'),[])
        self.assertEqual(c.dates_in_text('N 15 021 000021666'),[])
    def test_nivel_confianza(self):
        v=lambda vista,fechas:{'vista':vista,'fechas':fechas,'texto':''}
        F='2026-09-15'; G='2026-09-16'
        casos=[
            # (variantes, propuesta, nivel esperado, lectores que coinciden)
            ([v('trocr_color',[F]),v('vlm_color',[F])],F,'ALTA',['trocr','vlm']),
            ([v('color',[F]),v('trocr_color',[F]),v('vlm_color',[F])],F,'ALTA',['detector','trocr','vlm']),
            ([v('color',[F]),v('trocr_color',[F])],F,'MEDIA',['detector','trocr']),       # impreso + un manuscrito
            ([v('color',[F]),v('tinta',[F]),v('contraste',[F])],F,'MEDIA',['detector']),   # solo el detector impreso
            ([v('trocr_color',[F]),v('trocr_tinta',[F])],F,'MEDIA',['trocr']),             # un solo manuscrito, dos vistas
            ([v('vlm_color',[F])],F,'MEDIA',['vlm']),
            ([v('color',[F]),v('trocr_color',[F]),v('vlm_color',[])],F,'MEDIA',['detector','trocr']),
            ([v('trocr_color',[F]),v('vlm_color',[G])],None,'CONTRADICTORIA',[]),
            ([v('color',[G]),v('trocr_color',[F]),v('vlm_color',[F])],None,'CONTRADICTORIA',[]),  # el detector contradice
            ([v('trocr_color',[F]),v('vlm_color',[F]),v('acta_negativa',[G])],F,'ALTA',['trocr','vlm']),  # el acta no interviene
            ([v('color',[]),v('trocr_color',[])],None,'SIN_FECHA',[]),
            ([],None,'SIN_FECHA',[]),
        ]
        for variantes,prop,nivel,lectores in casos:
            with self.subTest(nivel=nivel,variantes=[x['vista'] for x in variantes]):
                self.assertEqual(c.nivel_confianza(variantes,prop),(nivel,lectores))
        class M: vlm=None
        self.assertIn('no es alcanzable',c.aviso_lectores(M()))
        M.vlm=object(); self.assertIn('Qwen2.5-VL',c.aviso_lectores(M()))
    def test_fallo_de_carga_del_vlm_no_detiene_el_motor(self):
        m=c.MotorOCR.__new__(c.MotorOCR)
        m.vlm=None; m.gpu=False; m.nombre='detector + trocr'
        def falla(model_name,gpu): raise RuntimeError('CUDA out of memory (simulado)')
        m._cargar_vlm=falla
        with contextlib.redirect_stdout(io.StringIO()) as out:
            m._activar_vlm('Qwen/Qwen2.5-VL-3B-Instruct')
        self.assertIsNone(m.vlm); self.assertIn('out of memory',m.vlm_error)
        self.assertIn('no se pudo cargar el lector VLM',out.getvalue()); self.assertEqual(m.nombre,'detector + trocr')
        self.assertIn('no se cargó',c.aviso_lectores(m)); self.assertIn('no es alcanzable',c.aviso_lectores(m))
    def test_alertas_plausibilidad(self):
        original=c.VENTANA_DIAS
        try:
            c.configurar(ventana_dias=30)
            t={'emision':'2026-09-02'}
            self.assertEqual(c.alertas({'fecha_notificacion':'15/09/2026','resolucion':'UNO','fecha_emision':'02/09/2026'},t),[])
            self.assertTrue(any('resoluci' in a for a in c.alertas({'fecha_notificacion':'15/09/2026'},t)))
            avisos=c.alertas({'fecha_notificacion':'15/12/2026','resolucion':'UNO','fecha_emision':'02/09/2026'},t)
            self.assertEqual(len(avisos),1); self.assertIn('104 días',avisos[0])
            self.assertTrue(any('domingo' in a for a in c.alertas({'fecha_notificacion':'06/09/2026','resolucion':'UNO','fecha_emision':'02/09/2026'},t)))
            # Cruce con la fecha del acta: solo alerta, nunca cambia la confianza.
            d={'fecha_notificacion':'22/09/2026','resolucion':'UNO','fecha_emision':'02/09/2026'}
            self.assertEqual(c.alertas(d,t,{'fecha_acta':['2026-09-22']}),[])
            av=c.alertas(d,t,{'fecha_acta':['2026-09-21']})
            self.assertEqual(len(av),1); self.assertIn('21/09/2026',av[0]); self.assertIn('DOMICILIO CERRADO',av[0])
            self.assertEqual(c.alertas(d,t,{'fecha_acta':[]}),[]); self.assertEqual(c.alertas(d,t,None),[])
            with self.assertRaises(ValueError): c.configurar(ventana_dias=0)
        finally:
            c.VENTANA_DIAS=original

class ExportacionLegalTests(ProyectoBase):
    def preparar(self):
        p=self.proyecto()
        with contextlib.redirect_stdout(io.StringIO()):
            p.inventario(); p.procesar(MotorFalso(),limite=10)
        return p
    def datos(self,n=1):
        return dict(expediente=f'16-012-{n}-2026',codigo='123',resolucion='UNO',fecha_notificacion='15/09/2026',fecha_emision='02/09/2026')
    def test_aprobaciones_evidencia_y_doble_revision(self):
        p=self.preparar(); rid=sorted(p.active)[0]
        p.records[rid]['confianza']='MEDIA'; p.records[rid]['evidencia_ocr']['roi_fecha']=[0.1,0.6,0.9,0.9]
        p.aprobar(rid,self.datos(),'Ana')
        r=p.records[rid]
        self.assertEqual(r['estado'],'APROBADO'); self.assertEqual(len(r['aprobaciones']),1)
        self.assertEqual(r['aprobacion']['confianza_ocr'],'MEDIA')
        self.assertTrue((self.salida/'evidencia'/f'{rid}_cargo.jpg').exists())
        self.assertTrue((self.salida/'evidencia'/f'{rid}_fecha.jpg').exists())
        with self.assertRaises(ValueError):  # misma persona, mismos datos
            p.aprobar(rid,self.datos(),'Ana')
        with contextlib.redirect_stdout(io.StringIO()):
            xlsx,csv_path=p.exportar(doble_revision='no_alta')
        self.assertIsNone(xlsx)
        with csv_path.open(encoding='utf-8-sig') as fh: filas=list(__import__('csv').DictReader(fh,delimiter=';'))
        fila=next(f for f in filas if f['id']==rid)
        self.assertEqual(fila['resultado'],'FALTA SEGUNDA REVISION'); self.assertEqual(fila['confianza_ocr'],'MEDIA')
        self.assertIn('_fecha.jpg',fila['evidencia'])
        p.aprobar(rid,self.datos(),'Luis')
        self.assertEqual(len({a['revisor'] for a in p.aprobaciones_vigentes(p.records[rid])}),2)
        self.assertEqual(p.lista_revision('segunda'),[])
        with contextlib.redirect_stdout(io.StringIO()):
            xlsx,csv_path=p.exportar(doble_revision='no_alta')
        self.assertIsNotNone(xlsx)
        actas=sorted(self.salida.glob('acta_exportacion_*.json'))
        acta=json.loads(actas[-1].read_text(encoding='utf-8'))
        self.assertEqual(acta['celdas_escritas'],1); self.assertEqual(acta['excel_generado']['sha256'],c.sha256(xlsx))
        self.assertEqual(acta['detalle'][0]['celda'],'F2'); self.assertEqual(acta['detalle'][0]['revisor'],'Ana | Luis')
        # Una corrección posterior invalida las aprobaciones anteriores.
        corregido=self.datos(); corregido['fecha_notificacion']='16/09/2026'
        p.aprobar(rid,corregido,'Ana')
        self.assertEqual(len(p.aprobaciones_vigentes(p.records[rid])),1)
        self.assertEqual(p.lista_revision('segunda'),[rid])
        with contextlib.redirect_stdout(io.StringIO()):
            xlsx,_=p.exportar(doble_revision=True)
        self.assertIsNone(xlsx)
        with contextlib.redirect_stdout(io.StringIO()):
            xlsx,_=p.exportar(doble_revision=False)
        self.assertIsNotNone(xlsx)
        self.assertEqual(load_workbook(xlsx)['2026']['F2'].value.date(),date(2026,9,16))
    def test_actualizar_excel_en_drive_con_respaldo(self):
        p=self.preparar(); rid=sorted(p.active)[0]
        drive=Path(self.temp.name)/'PROYECTO REC 1'/'expedientes.xlsx'; drive.parent.mkdir()
        drive.write_bytes(self.excel.read_bytes())
        p.aprobar(rid,self.datos(),'Ana')
        with contextlib.redirect_stdout(io.StringIO()):
            copia,_=p.exportar()
        huella_original=p.fingerprint
        with contextlib.redirect_stdout(io.StringIO()):
            destino,respaldo=p.actualizar_excel_en_drive(copia,drive)
        self.assertEqual(destino,drive); self.assertTrue(respaldo.exists()); self.assertIn('_respaldo_',respaldo.name)
        self.assertEqual(c.sha256(respaldo),huella_original)            # respaldo = original intacto
        self.assertEqual(c.sha256(drive),c.sha256(copia))               # Drive = copia verificada
        self.assertEqual(c.sha256(p.excel),c.sha256(drive))             # copia de entrada renovada
        self.assertEqual(p.fingerprint,c.sha256(drive))
        cfg=json.loads((self.salida/'proyecto.json').read_text())
        self.assertEqual(cfg['excel_sha256'],p.fingerprint); self.assertEqual(cfg['excel_sha256_historial'][0]['anterior'],huella_original)
        self.assertEqual(load_workbook(drive)['2026']['F2'].value.date(),date(2026,9,15))
        # El índice ya ve F2 registrada: una nueva aprobación sobre esa celda se rechaza y una exportación no la reescribe.
        with self.assertRaises(ValueError):
            p.aprobar(rid,self.datos(),'Luis')
        with contextlib.redirect_stdout(io.StringIO()):
            xlsx,_=p.exportar()
        self.assertIsNone(xlsx)
        # Reabrir la carpeta de resultados con el Excel actualizado es válido; con el respaldo (huella anterior) no.
        with contextlib.redirect_stdout(io.StringIO()):
            c.Proyecto(self.fotos,drive,self.salida)
        with self.assertRaises(ValueError):
            c.Proyecto(self.fotos,respaldo,self.salida)
        # Si el Excel de Drive cambió a mano, no se sobrescribe.
        drive.write_bytes(drive.read_bytes()+b' ')
        with self.assertRaises(ValueError):
            p.actualizar_excel_en_drive(copia,drive)
    def test_resolucion_no_leida_exige_casilla(self):
        p=self.preparar(); rid=sorted(p.active)[0]
        sin=self.datos(); sin['resolucion']=None
        with self.assertRaises(ValueError) as cm:
            p.aprobar(rid,sin,'Ana')
        self.assertIn('REC UNO',str(cm.exception)); self.assertEqual(p.records[rid]['estado'],'POR_REVISAR')
        otra=self.datos(); otra['resolucion']='DOS'
        with self.assertRaises(ValueError):
            p.aprobar(rid,otra,'Ana',verifico_rec_uno=True)
        p.aprobar(rid,sin,'Ana',verifico_rec_uno=True)
        self.assertTrue(p.records[rid]['aprobacion']['verifico_rec_uno'])
        self.assertTrue(any('resoluci' in a for a in p.records[rid]['aprobacion']['alertas']))
        p.aprobar(rid,self.datos(),'Luis')  # con resolución leída no hace falta la casilla
        self.assertFalse(p.records[rid]['aprobacion']['verifico_rec_uno'])
        with contextlib.redirect_stdout(io.StringIO()):
            xlsx,csv_path=p.exportar()
        with csv_path.open(encoding='utf-8-sig') as fh: filas=list(__import__('csv').DictReader(fh,delimiter=';'))
        self.assertEqual(next(f for f in filas if f['id']==rid)['verifico_rec_uno'],'no')
        actas=sorted(self.salida.glob('acta_exportacion_*.json'))
        self.assertIn('verifico_rec_uno',json.loads(actas[-1].read_text(encoding='utf-8'))['detalle'][0])

def caja(x1,y1,x2,y2,text,score=0.9):
    return {'box':[[x1,y1],[x2,y1],[x2,y2],[x1,y2]],'text':text,'score':score}

class LocalizacionTests(unittest.TestCase):
    W,H=2200,2800
    def test_rotulos_mal_leidos_por_easyocr(self):
        rows=[caja(1000,1900,1500,1930,'FECHA DE HotECACLON:'),caja(1000,1990,1500,2020,'HORADENOIIECACKN:'),
              caja(1000,2080,1500,2110,'FOLIOS ENTREGADOS:'),caja(100,1900,600,1930,'CARACTERISTICAS DEL LUGAR')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulos_fecha_y_hora')
        self.assertEqual((roi[1],roi[3]),(1921,1996)); self.assertLess(roi[0],1000); self.assertGreater(roi[2],1500)
    def test_rotulo_partido_en_dos_detecciones(self):
        rows=[caja(1000,1900,1180,1930,'FECHA DE'),caja(1190,1901,1500,1931,'NOTIFICACION:'),caja(1000,1990,1500,2020,'HORA DE NOTIFICACION:')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulos_fecha_y_hora'); self.assertEqual((roi[1],roi[3]),(1921,1999))
    def test_solo_hora_o_nada(self):
        roi,metodo=c.localizar_recuadro_fecha([caja(1000,1990,1500,2020,'HORA DE NOTIFICACION:')],self.W,self.H)
        self.assertEqual(metodo,'rotulo_hora'); self.assertEqual(roi[3],1990); self.assertLess(roi[1],1990)
        roi,metodo=c.localizar_recuadro_fecha([caja(1000,1990,1500,2020,'FECHA Y HORA'),caja(1000,1890,1500,1920,'FOLIOS ENTREGADOS')],self.W,self.H)
        self.assertIsNone(roi)
        # 'HORA DE NOTIFICACION' no debe aceptarse como rótulo FECHA aunque se parezca.
        roi,metodo=c.localizar_recuadro_fecha([caja(1000,1990,1500,2020,'HORA DE NOTIFICACION')],self.W,self.H)
        self.assertEqual(metodo,'rotulo_hora')
    def test_inicio_mal_leido_y_rotulos_de_otras_fechas(self):
        # Inicio mal leído pero NOTIFICACION reconocible.
        rows=[caja(1000,1900,1500,1930,'PECHA DE NOTIFICACION:'),caja(1000,1990,1500,2020,'HORA DE NOTIFICACION:')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulos_fecha_y_hora'); self.assertEqual((roi[1],roi[3]),(1921,1996))
        # Un rótulo HORA pegado al de FECHA ya no deja un recorte de 30 px: altura mínima de 2,5 rótulos.
        rows=[caja(1000,1900,1500,1930,'FECHA DE NOTIFICACION:'),caja(1000,1945,1500,1975,'HORA DE NOTIFICACION:')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulos_fecha_y_hora'); self.assertGreaterEqual(roi[3]-roi[1],75)
        # Solo la palabra NOTIFICACION, con el rótulo HORA debajo: es el rótulo de FECHA.
        rows=[caja(1000,1900,1500,1930,'NOTIFICACION:'),caja(1000,1990,1500,2020,'HORA DE NOTIFICACION:')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulo_notificacion_fecha'); self.assertEqual((roi[1],roi[3]),(1921,1996))
        # Dos rótulos leídos solo como NOTIFICACION: el de arriba es FECHA y el de abajo HORA.
        rows=[caja(1000,1900,1500,1930,'NOTIFICACION:'),caja(1000,1990,1500,2020,'NOTIFICACION:')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulo_notificacion_fecha'); self.assertEqual((roi[1],roi[3]),(1921,1996))
        # 'NOTIFICACIÓN' dentro de un párrafo (línea larga) o bajo 'TÉNGASE POR BIEN NOTIFICADO' no es ancla.
        parrafo=caja(100,2300,1500,2330,'Se adjunta a la presente copia de la O.P. y su cargo de notificacion y constancia de consentimiento')
        self.assertEqual(c.localizar_recuadros_fecha([parrafo],self.W,self.H)[0],[])
        bajo_cierre=[caja(100,2200,1500,2230,'TENGASE POR BIEN NOTIFICADO, SEGUN LO DISPUESTO'),caja(1000,2250,1200,2280,'NOTIFICACION:')]
        self.assertEqual(c.localizar_recuadros_fecha(bajo_cierre,self.W,self.H)[0],[])
        # Entre varias anclas ambiguas gana la que tiene HORA o FOLIOS debajo, aunque esté más arriba.
        rows=[caja(1000,1700,1200,1730,'NOTIFICACION:'),caja(1000,1900,1200,1930,'NOTIFICACION:'),caja(1000,1990,1500,2020,'HORA DE NOTIFICACION:'),
              caja(1000,2100,1200,2130,'NOTIFICACION')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulo_notificacion_fecha'); self.assertEqual(roi[1],1921)
        # Sin apoyo debajo, gana la más baja de la página (la tabla está al final).
        rows=[caja(1000,1700,1200,1730,'NOTIFICACION:'),caja(1000,2100,1200,2130,'NOTIFICACION')]
        recs,_=c.localizar_recuadros_fecha(rows,self.W,self.H)
        self.assertEqual(recs[0][0][1],2130)
        # NOTIFICACION suelto con FOLIOS debajo: es el rótulo de HORA, la casilla está encima.
        rows=[caja(1000,1900,1500,1930,'NOTIFICACION:'),caja(1000,1990,1500,2020,'FOLIOS ENTREGADOS:')]
        roi,metodo=c.localizar_recuadro_fecha(rows,self.W,self.H)
        self.assertEqual(metodo,'rotulo_notificacion_hora'); self.assertEqual(roi[3],1900); self.assertLess(roi[1],1900)
        # NOTIFICACION suelto y nada más: se usa como ancla con las dos casillas (encima y debajo).
        recs,motivo=c.localizar_recuadros_fecha([caja(1000,1900,1500,1930,'NOTIFICACION:')],self.W,self.H)
        self.assertEqual([m for _,m in recs],['rotulo_notificacion_ambos']*2)
        self.assertEqual(recs[0][0][1],1930); self.assertEqual(recs[1][0][3],1900)
        # NOTIFICACION suelto en la mitad superior de la página no es ancla; junto a EMISIÓN/VENCIMIENTO tampoco.
        self.assertEqual(c.localizar_recuadros_fecha([caja(1000,300,1500,330,'NOTIFICACION:')],self.W,self.H)[0],[])
        for texto in ('NOTIFICACION DE EMISION','FECHA DE NOTIFICACION DE VENCIMIENTO'):
            self.assertEqual(c.localizar_recuadros_fecha([caja(1000,1900,1500,1930,texto)],self.W,self.H)[0],[])
        # Rótulos de otras fechas no se aceptan como ancla.
        for texto in ('FECHA DE EMISION','FECHA DE VENCIMIENTO','FECHA DE EMISION REC','FECHA DE EMISION:','FECHA Y HORA','FECHA:'):
            with self.subTest(texto=texto):
                self.assertEqual(c.parecido_rotulo(texto,'fecha','hora'),0.0)
                self.assertIsNone(c.localizar_recuadro_fecha([caja(1000,1900,1500,1930,texto)],self.W,self.H)[0])
        # Positivos del informe de campo y negativos cruzados.
        self.assertGreater(c.parecido_rotulo('FECHA DE HotECACLON:','fecha','hora'),0.65)
        self.assertGreater(c.parecido_rotulo('HORADENOIIECACKN:','hora','fecha'),0.65)
        self.assertEqual(c.parecido_rotulo('HORA DE NOTIFICACION','fecha','hora'),0.0)
        self.assertEqual(c.parecido_rotulo('FECHA DE NOTIFICACION','hora','fecha'),0.0)
        self.assertEqual(c.parecido_rotulo('HORADENOIIECACKN:','fecha','hora'),0.0)
    def test_acta_negativa_y_clasificacion(self):
        rows=[caja(100,1700,900,1730,'Siendo las 9:00 horas, del dia 11/09/2026 , quien suscribe'),
              caja(1000,1900,1500,1930,'FECHA DE NOTIFICACION:'),caja(1000,1990,1500,2020,'HORA DE NOTIFICACION:')]
        self.assertEqual(c.fecha_acta_negativa(rows),['2026-09-11'])
        head='EXPEDIENTE COACTIVO : 16-012-000000001-2026-SAT-H/DCC-IPA\nRESOLUCION DE EJECUCION COACTIVA'
        self.assertEqual(c.clasificar_documento(head,rows,self.W,self.H,'16-012-000000001-2026'),'REC')
        self.assertEqual(c.clasificar_documento(head,[],2480,1748,'16-012-000000001-2026'),'REC_SIN_CARGO')      # horizontal: media página
        self.assertEqual(c.clasificar_documento(head,[],self.W,self.H,'16-012-000000001-2026'),'REC_CARGO_NO_UBICADO')  # vertical: rótulo no leído
        self.assertEqual(c.clasificar_documento('texto suelto',rows,self.W,self.H,None),'CARGO_SIN_ENCABEZADO')
        self.assertEqual(c.clasificar_documento('AVISO PREVIO DE NOTIFICACION',[],3307,2330,None),'ANEXO')
        self.assertEqual(c.clasificar_documento('CEDULON',[],3307,2330,None),'ANEXO')
        self.assertEqual(c.clasificar_documento('nada',[],self.W,self.H,None),'DESCONOCIDO')
    def test_nombre_y_emision_impresa(self):
        txt=('SERVICIO DE ADMINISTRACION TRIBUTARIA\nEXPEDIENTE COACTIVO : 16-012-000000001-2026-SAT-H/DCC-IPA\n'
             'OBLIGADO (A) : JUANA FICTICIA PEREZ Y PEDRO FICTICIO LOPEZ\nDOMICILIO FISCAL : ASOC. EJEMPLO\n'
             'CODIGO CONTRIBUYENTE: 000000000123\nResolución N° UNO\nAyacucho, 02 de setiembre de 2026')
        self.assertEqual(c.parse_header_extra(txt),('JUANA FICTICIA PEREZ Y PEDRO FICTICIO LOPEZ','2026-09-02'))
        # Rótulo mal leído, nombre pegado al siguiente rótulo y mes con errores de OCR.
        self.assertEqual(c.parse_header_extra('OBLlGADO (A): MARIA FICTICIA QUISPE DOMICILIO FISCAL: X\nAyacucho, 02 de septiernbre de 2026'),
                         ('MARIA FICTICIA QUISPE','2026-09-02'))
        self.assertEqual(c.fecha_en_letras('AYACUCHO, 2 DE OCTUBRE DEL 2026'),'2026-10-02')
        self.assertIsNone(c.fecha_en_letras('2 DE XXXXXX DE 2026')); self.assertIsNone(c.fecha_en_letras('sin fecha'))
        self.assertEqual(c.parse_header_extra('texto sin campos'),(None,None))
        # Similitud de nombres por palabras, tolerante al orden y a un segundo obligado.
        self.assertGreaterEqual(c.similitud_nombres('JUANA FICTICIA PEREZ Y PEDRO FICTICIO LOPEZ','PEREZ JUANA FICTICIA'),0.5)
        self.assertEqual(c.similitud_nombres('JUANA PEREZ','CARLOS RAMOS'),0.0)
        # La emisión impresa distinta de la columna E bloquea; el nombre distinto solo alerta.
        src=Path(tempfile.mkdtemp())/'o.xlsx'; libro_ficticio(src); index,_,_=c.build_index(src,('2026',))
        d=dict(expediente='16-012-1-2026',codigo='123',resolucion='UNO',nombre='PERSONA FICTICIA',fecha_notificacion='15/09/2026',fecha_emision='03/09/2026')
        target,issues=c.evaluar(d,index); self.assertTrue(any('emisión impresa' in i for i in issues))
        d['fecha_emision']='02/09/2026'; target,issues=c.evaluar(d,index); self.assertEqual(issues,[])
        self.assertEqual(c.alertas(d,target),[])
        d['nombre']='OTRO NOMBRE DISTINTO'; self.assertTrue(any('nombre leído' in a for a in c.alertas(d,target)))
        d['nombre']=None; d['fecha_emision']=None
        self.assertTrue(any('emisión impresa' in a for a in c.alertas(d,target)))
    def test_parse_header_resolucion_tolerante(self):
        exp,code,res=c.parse_header('EXPEDIENTE COACTIVO : 16-012-000000001-2026-SAT-H/DCC-IPA\nCODIGO CONTRIBUYENTE: 000000000123\nResoluci6n N° UNO')
        self.assertEqual((exp,code,res),('16-012-000000001-2026-SAT-H/DCC-IPA','123','UNO'))
        self.assertEqual(c.parse_header('RESOLUCLON N UNO')[2],'UNO')
        self.assertIsNone(c.parse_header('texto sin resolucion')[2])
        d=dict(expediente='16-012-1-2026',codigo='123',resolucion=None,fecha_notificacion='15/09/2026')
        src=Path(tempfile.mkdtemp())/'o.xlsx'; libro_ficticio(src); index,_,_=c.build_index(src,('2026',))
        self.assertEqual(c.evaluar(d,index)[1],[])
        d['resolucion']='DOS'; self.assertTrue(any('REC UNO' in i for i in c.evaluar(d,index)[1]))

class MotorMediaPagina(MotorFalso):
    """Simula una REC partida en dos fotos: encabezado en una, cargo en la otra."""
    def leer(self,src):
        with Image.open(src) as im: color=im.getpixel((0,0))
        base=dict(expediente=None,codigo=None,resolucion=None,fecha_notificacion=None,fecha_emision=None,fecha_texto=None,observaciones='')
        if color[0]==0:   # f1: encabezado sin cargo
            return dict(base,expediente='16-012-1-2026',codigo='123',resolucion='UNO'),{'motor':self.nombre,'tipo':'REC_SIN_CARGO'}
        if color[0]==40:  # f2: cargo sin encabezado
            return dict(base,fecha_notificacion='15/09/2026'),{'motor':self.nombre,'tipo':'CARGO_SIN_ENCABEZADO'}
        return dict(base),{'motor':self.nombre,'tipo':'ANEXO'}

class EmparejamientoTests(ProyectoBase):
    def test_emparejar_y_medir(self):
        p=self.proyecto()
        with contextlib.redirect_stdout(io.StringIO()):
            p.inventario(); resumen=p.procesar(MotorMediaPagina(),limite=10)
        self.assertEqual(resumen['emparejadas'],0)  # f1 y f2 están en carpetas distintas
        f1=p.ids_de([self.fotos/'A/sub/f1.png'])[0]; f2=p.ids_de([self.fotos/'A/f2.png'])[0]; f3=p.ids_de([self.fotos/'B/f3.png'])[0]
        self.assertEqual(p.records[f3]['estado'],'ANEXO')
        self.assertEqual(sorted(p.lista_revision('pendientes')),sorted([f1,f2])); self.assertIn(f3,p.lista_revision('todos'))
        # f1 y f2 están en carpetas distintas: no se emparejan. Movemos f2 junto a f1 y repetimos.
        self.assertIsNone(p.records[f2]['pareja'])
        p.records[f1]['archivos']=['A/Imagen (149).png']; p.records[f2]['archivos']=['A/Imagen (150).png']
        self.assertEqual(p.emparejar(),1)
        self.assertEqual(p.records[f2]['pareja'],f1); self.assertEqual(p.records[f1]['pareja'],f2)
        self.assertEqual(p.records[f2]['lectura']['expediente'],'16-012-1-2026')
        self.assertEqual(p.records[f2]['destino']['celda'],'F2'); self.assertEqual(p.records[f2]['problemas'],[])
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(p.resumen_alcance(),{'filas':5,'f_vacia':3,'f_registrada':2})
        # Medición contra F ya registrada: la fila 3 (16-012-3) tiene 10/09/2026.
        p.records[f2]['lectura'].update(expediente='16-012-3-2026',fecha_notificacion='10/09/2026')
        p.records[f1]['lectura'].update(fecha_notificacion='11/09/2026',expediente='16-012-3-2026')
        p.records[f2]['confianza']='ALTA'; p.records[f1]['confianza']='MEDIA'
        with contextlib.redirect_stdout(io.StringIO()) as out:
            res=p.medir_acierto()
        self.assertEqual((res['con_referencia'],res['aciertos'],res['errores']),(2,1,1))
        self.assertEqual(res['tabla']['ALTA'],{'total':1,'ACIERTO':{'n':1,'pct':100.0},'ERROR':{'n':0,'pct':0.0},'SIN_FECHA':{'n':0,'pct':0.0}})
        self.assertEqual(res['tabla']['MEDIA']['ERROR'],{'n':1,'pct':100.0}); self.assertEqual(res['tabla']['TOTAL']['total'],2)
        self.assertFalse(res['concluyente']); self.assertIn('no es concluyente',out.getvalue())
        self.assertIn('confianza',out.getvalue()); self.assertIn('ACIERTO',out.getvalue())
        guardado=json.loads(Path(res['archivo']).read_text(encoding='utf-8'))
        self.assertEqual(guardado['tabla']['TOTAL']['total'],2); self.assertTrue(Path(res['archivo']).name.startswith('piloto_'))
        self.assertIn('tipos',guardado); self.assertIn('confianza_muestra',guardado)
    def test_sondeos_y_conciliacion(self):
        p=self.proyecto()
        with contextlib.redirect_stdout(io.StringIO()) as out:
            p.inventario(); fmt=p.sondeo_formato()
        self.assertEqual(fmt['fotos'],4); self.assertEqual(fmt['por_tamano'],{'40x30':4}); self.assertEqual(fmt['errores'],[])
        self.assertEqual(fmt['por_categoria_y_orientacion'],{'OTRA horizontal':4}); self.assertTrue(Path(fmt['archivo']).exists())
        self.assertEqual(c.Proyecto.categoria_carpeta('RECs/REC PARA/REC/NEGATIVA/Imagen (1).jpg'),'NEGATIVA')
        self.assertEqual(c.Proyecto.categoria_carpeta('X/DOMICILIO CERRADO/a.jpg'),'DOMICILIO_CERRADO')
        self.assertEqual(c.Proyecto.categoria_carpeta('X/AVISO PREVIO/a.jpg'),'AVISO_PREVIO')
        with contextlib.redirect_stdout(io.StringIO()) as out:
            antes=p.conciliar(guardar=False)
        self.assertEqual((antes['filas_en_alcance'],antes['filas_f_vacia'],antes['fotos_archivos'],antes['fotos_unicas'],antes['fotos_no_leidas']),(3,1,4,3,3))
        self.assertIn('parcial',out.getvalue())
        with contextlib.redirect_stdout(io.StringIO()):
            p.procesar(MotorFalso(),limite=10); est=p.sondeo_estructura(); con=p.conciliar()
        self.assertEqual(est['fotos_leidas'],3); self.assertEqual(est['estructura_completa'],3); self.assertEqual(est['por_tipo'],{'REC':3})
        self.assertEqual(est['campos_encontrados']['fecha_notificacion'],3); self.assertTrue(Path(est['archivo']).exists())
        self.assertEqual((con['filas_con_foto'],con['filas_sin_foto'],con['filas_con_varias_fotos'],con['expedientes_fuera_del_lote'],con['fotos_sin_expediente'],con['fotos_no_leidas']),(1,2,1,0,0,0))
        self.assertEqual([x['celda'] for x in con['detalle']['filas_sin_foto']],['F4','F5'])
        self.assertEqual(json.loads(Path(con['archivo']).read_text(encoding='utf-8'))['filas_en_alcance'],3)
        # Un expediente leído que no está en el lote se reporta como fuera del lote.
        k=sorted(p.active)[0]; p.records[k]['lectura']['expediente']='16-012-2-2026'
        with contextlib.redirect_stdout(io.StringIO()):
            con=p.conciliar(guardar=False)
        self.assertEqual(con['expedientes_fuera_del_lote'],1)

class PaddleWorkerTests(unittest.TestCase):
    """El trabajador se prueba en modo simulado (sin modelos): protocolo, caída, reinicio y relevo a EasyOCR."""
    FAKE={'PADDLE_WORKER_FAKE':'1'}
    def test_protocolo_caida_y_reinicio(self):
        w=c.PaddleWorker(sys.executable,{'lang':'es'},env=self.FAKE)
        try:
            self.assertTrue(w.info['ready']); self.assertTrue(w.info['fake']); self.assertTrue(w.vivo())
            rows=w.read(Image.new('RGB',(20,20),'white'))
            self.assertEqual(len(rows),1); self.assertTrue(rows[0]['text'].startswith('FAKE')); self.assertEqual(len(rows[0]['box']),4)
            w._enviar({'cmd':'crash'})
            with self.assertRaises(c.PaddleWorkerError) as cm:
                w.read(Image.new('RGB',(20,20),'white'))
            self.assertIn('139',str(cm.exception)); self.assertFalse(w.vivo())
            w.start(); self.assertTrue(w.vivo()); self.assertEqual(w.arranques,2)
            self.assertEqual(len(w.read(Image.new('RGB',(20,20),'white'))),1)
        finally:
            w.close()
        self.assertFalse(w.vivo())
    def test_arranque_fallido_se_informa(self):
        # Un intérprete sin paddleocr no cuelga: responde con error y PaddleWorker lo convierte en excepción.
        with self.assertRaises(c.PaddleWorkerError) as cm:
            c.PaddleWorker(sys.executable,{'lang':'es'},env={'PADDLE_WORKER_FAKE':'','PYTHONPATH':'/nonexistent','PYTHONNOUSERSITE':'1','PYTHONSAFEPATH':'1'},timeout_inicio=60) if not _paddle_disponible() else _simular_fallo()
        self.assertIn('PaddleOCR no pudo iniciarse',str(cm.exception))
    def test_probar_paddle_con_interprete_inexistente(self):
        ok,motivo=c.probar_paddle(python='/ruta/inexistente/python',kwargs={'lang':'es'})
        self.assertFalse(ok); self.assertIn('no se pudo ejecutar',motivo)
    def test_motor_releva_a_easyocr_en_la_primera_caida(self):
        m=c.MotorOCR.__new__(c.MotorOCR)
        m.paddle=None; m.reader=None; m.caidas_paddle=0; m.gpu=False; m.nombre='PaddleOCR PP-OCRv5 (cpu, proceso aislado) + trocr'
        m.paddle_worker=c.PaddleWorker(sys.executable,{'lang':'es'},env=self.FAKE)
        activaciones=[]
        def falso_easyocr():
            activaciones.append(1); m.detector_nombre='EasyOCR (cpu)'
            class R:
                def readtext(self,arr,**kw): return [([[0,0],[5,0],[5,5],[0,5]],'EASY',0.9)]
            m.reader=R()
        m._activar_easyocr=falso_easyocr
        im=Image.new('RGB',(20,20),'white')
        self.assertTrue(m.read(im)[0]['text'].startswith('FAKE'))
        m.paddle_worker._enviar({'cmd':'crash'})
        with contextlib.redirect_stdout(io.StringIO()) as out:
            rows=m.read(im)                       # primera caída: relevo inmediato, sin reintentar Paddle
        self.assertEqual(rows[0]['text'],'EASY'); self.assertIsNone(m.paddle_worker); self.assertEqual(activaciones,[1])
        self.assertEqual(m.caidas_paddle,1); self.assertIn('Se pasa a EasyOCR',out.getvalue()); self.assertNotIn('Se reinicia',out.getvalue())
        self.assertTrue(m.nombre.startswith('EasyOCR (cpu) + trocr'))
        self.assertEqual(m.read(im)[0]['text'],'EASY')
    def test_kwargs_del_aislado(self):
        m=c.MotorOCR.__new__(c.MotorOCR); m.paddle_kwargs=c.kwargs_paddle(); m.paddle_kwargs['device']='cpu'
        kw=m.paddle_kwargs_aislado()
        self.assertTrue(kw['enable_mkldnn']); self.assertEqual((kw['text_det_limit_side_len'],kw['text_det_limit_type']),(1600,'max'))
        self.assertFalse(m.paddle_kwargs['enable_mkldnn'])

def _paddle_disponible():
    try:
        import paddleocr  # noqa: F401
        return True
    except Exception:
        return False

def _simular_fallo():
    # Con paddleocr instalado, forzamos un error de arranque pasando un argumento inválido al constructor.
    return c.PaddleWorker(sys.executable,{'argumento_invalido':1},env={'PADDLE_WORKER_FAKE':''},timeout_inicio=120)

def pagina_sintetica(con_rotulo=True, tinta=(40,60,200), fecha='15/09/2026'):
    """Página REC vertical de 1000x1400: encabezado impreso (negro) y fecha manuscrita (azul) en la zona del cargo."""
    from PIL import ImageFont
    im=Image.new('RGB',(1000,1400),'white'); dr=ImageDraw.Draw(im)
    try: fuente=ImageFont.load_default(size=34)
    except TypeError: fuente=ImageFont.load_default()
    dr.text((60,60),'EXPEDIENTE COACTIVO : 16-012-000000001-2026-SAT-H/DCC-IPA',fill='black',font=fuente)
    dr.line((60,950,940,950),fill='black',width=2)
    if con_rotulo:
        dr.text((480,958),'FECHA DE NOTIFICACION:',fill='black',font=fuente)
    dr.text((500,1000),fecha,fill=tinta,font=fuente)
    dr.line((60,1060,940,1060),fill='black',width=2)
    return im

class MotorSintetico:
    """Motor con lectores simulados sobre imágenes sintéticas.

    read(): devuelve rótulos solo para el pie de página (según la variante) y nada para los recortes.
    TrOCR y el VLM 'leen' la fecha únicamente cuando el recorte contiene tinta azul.
    """
    def __init__(self, filas_pie, con_vlm=False, fecha='15/09/2026', vlm_localizador=None):
        self.filas_pie=filas_pie; self.fecha=fecha; self.nombre='sintetico'
        self.paddle=None; self.paddle_worker=None; self.reader=None; self.trocr=object()
        self.vlm=object() if con_vlm else None; self.vlm_localizador=vlm_localizador; self.llamadas_vlm=[]; self.vlm_vistas=2; self.api=None
    def read(self, im, digits=False):
        if im.size==(1000,448): return []           # encabezado: parse_header no necesita filas, usamos el texto
        if im.size==(1000,700): return self.filas_pie
        return []
    def _ve_tinta(self, im): return c.vista_tinta(im)[0] is not None
    def read_handwritten_lote(self, imgs): return [self.fecha if self._ve_tinta(im) else 'nada' for im in imgs]
    def read_vlm(self, im, prompt=None, max_pixels=None):
        self.llamadas_vlm.append((im.size,prompt is not None))
        if prompt is not None: return self.vlm_localizador or 'ILEGIBLE'
        return self.fecha if self._ve_tinta(im) else 'ILEGIBLE'
    leer=c.MotorOCR.leer; VLM_PROMPT_LOCALIZADOR=c.MotorOCR.VLM_PROMPT_LOCALIZADOR

class ImagenesSinteticasTests(unittest.TestCase):
    def guardar(self, im):
        if not hasattr(self,'temps'): self.temps=[]
        d=tempfile.TemporaryDirectory(); self.temps.append(d); ruta=Path(d.name)/'pagina.png'; im.save(ruta); return ruta
    def tearDown(self):
        for d in getattr(self,'temps',[]): d.cleanup()
    def encabezado(self):
        # parse_header() recibe el texto de las detecciones del encabezado; aquí lo simulamos con una fila.
        return [caja(60,60,900,100,'EXPEDIENTE COACTIVO : 16-012-000000001-2026-SAT-H/DCC-IPA')]
    def motor(self, filas_pie, **kw):
        m=MotorSintetico(filas_pie,**kw)
        cabecera=self.encabezado()
        original_read=m.read
        m.read=lambda im,digits=False: cabecera if im.size==(1000,448) else original_read(im,digits)
        return m
    def test_1_rotulo_ambiguo_como_ancla(self):
        ruta=self.guardar(pagina_sintetica())
        # El pie se recorta desde el 50 % (y=700): el rótulo 'NOTIFICACION:' queda en y 258-290 del recorte.
        pie=[caja(480,258,760,290,'NOTIFICACION:')]
        d,ev=self.motor(pie).leer(ruta)
        self.assertEqual(ev['tipo'],'REC'); self.assertEqual([r['metodo'] for r in ev['recuadros']],['rotulo_notificacion_ambos'])
        self.assertEqual(d['fecha_notificacion'],'2026-09-15'); self.assertEqual(ev['confianza'],'MEDIA')
        self.assertNotIn('recuadro_por_posicion',ev)
        vistas={v['vista']:v['fechas'] for v in ev['variantes_fecha']}
        self.assertEqual(vistas['trocr_color'],['2026-09-15']); self.assertNotIn('r2_trocr_color',vistas)   # la segunda casilla no se lee si la primera dio fecha
        # Si la primera casilla (debajo) no tiene la fecha, se lee la alternativa (encima).
        ruta_arriba=self.guardar(pagina_sintetica(con_rotulo=False)); im=Image.open(ruta_arriba); dr=ImageDraw.Draw(im)
        dr.rectangle((470,985,790,1080),fill='white'); dr.rectangle((480,900,760,935),fill='white'); im.save(ruta_arriba)
        from PIL import ImageFont
        try: fuente=ImageFont.load_default(size=34)
        except TypeError: fuente=ImageFont.load_default()
        dr.text((500,905),'15/09/2026',fill=(40,60,200),font=fuente); im.save(ruta_arriba)
        pie_hora=[caja(480,300,760,332,'NOTIFICACION:')]   # el rótulo detectado es en realidad el de HORA (y=1000-1032 en página)
        d,ev=self.motor(pie_hora).leer(ruta_arriba)
        self.assertEqual([r['metodo'] for r in ev['recuadros']],['rotulo_notificacion_ambos']*2)
        # Con el VLM coincidiendo sobre el recuadro anclado por rótulo, ALTA sí es alcanzable.
        d,ev=self.motor(pie,con_vlm=True).leer(ruta)
        self.assertEqual((d['fecha_notificacion'],ev['confianza']),('2026-09-15','ALTA'))
    def test_2_respaldo_por_posicion_nunca_alta(self):
        ruta=self.guardar(pagina_sintetica(con_rotulo=False))
        m=self.motor([],con_vlm=True)
        d,ev=m.leer(ruta)
        self.assertEqual(ev['tipo'],'REC_CARGO_NO_UBICADO'); self.assertEqual(ev['recuadros'],[])
        self.assertIn('recuadro_por_posicion',ev); self.assertIn('Respaldo',ev['metodo_recuadro']); self.assertGreaterEqual(ev['bandas_tinta'],1)
        self.assertEqual(d['fecha_notificacion'],'2026-09-15'); self.assertIn('recorte por posición',d['observaciones'])
        vistas={v['vista']:v for v in ev['variantes_fecha']}
        self.assertTrue(vistas['pos1_trocr_color']['por_posicion']); self.assertEqual(vistas['pos1_vlm_color']['fechas'],['2026-09-15'])
        self.assertIn('pos1_vlm_contraste',vistas); self.assertNotIn('pos1_vlm_tinta',vistas); self.assertNotIn('pos1_trocr_contraste',vistas)
        self.assertEqual(ev['confianza'],'MEDIA')                       # dos lectores manuscritos, pero por posición: nunca ALTA
        self.assertFalse(any(p for _,p in m.llamadas_vlm))             # con fecha leída no hace falta el localizador
        self.assertEqual(ev['roi_fecha'],ev['recuadro_por_posicion'])
    def test_3_vlm_como_localizador(self):
        # Fecha en tinta negra: no hay bandas de tinta azul, así que se pide al VLM que localice el valor.
        ruta=self.guardar(pagina_sintetica(con_rotulo=False,tinta=(0,0,0)))
        m=self.motor([],con_vlm=True,vlm_localizador='16/09/2026')
        d,ev=m.leer(ruta)
        self.assertEqual(ev['bandas_tinta'],0)
        self.assertEqual(d['fecha_notificacion'],'2026-09-16'); self.assertEqual(ev['confianza'],'MEDIA')
        v=[x for x in ev['variantes_fecha'] if x['vista']=='vlm_mitad_inferior'][0]
        self.assertTrue(v['por_posicion']); self.assertEqual(v['fechas'],['2026-09-16'])
        self.assertEqual([p for _,p in m.llamadas_vlm],[True]); self.assertEqual(m.llamadas_vlm[0][0],(1000,770))
        # Si el localizador responde ILEGIBLE, queda sin fecha y con la observación de completar a mano.
        m=self.motor([],con_vlm=True,vlm_localizador=None)
        d,ev=m.leer(ruta)
        self.assertIsNone(d['fecha_notificacion']); self.assertEqual(ev['confianza'],'SIN_FECHA'); self.assertIn('recorte por posición',d['observaciones'])
    def test_4_clasificacion_y_media_pagina(self):
        # Vertical con encabezado y sin rótulo: REC_CARGO_NO_UBICADO (se intenta por posición). Horizontal: REC_SIN_CARGO, sin respaldo.
        ruta=self.guardar(pagina_sintetica(con_rotulo=False))
        d,ev=self.motor([]).leer(ruta); self.assertEqual(ev['tipo'],'REC_CARGO_NO_UBICADO'); self.assertIn('recuadro_por_posicion',ev)
        apaisada=pagina_sintetica(con_rotulo=False).rotate(90,expand=True)
        ruta2=self.guardar(apaisada)
        m=self.motor([]); m.read=lambda im,digits=False: self.encabezado()
        d,ev=m.leer(ruta2)
        self.assertEqual(ev['tipo'],'REC_SIN_CARGO'); self.assertNotIn('recuadro_por_posicion',ev); self.assertIsNone(d['fecha_notificacion'])
        self.assertIn('Media página',d['observaciones'])
    def test_fechas_implausibles_no_se_proponen(self):
        # Qwen 'lee' 06/09/2000 en color y 15/09/2026 en contraste: el año 2000 se descarta y no crea contradicción.
        ruta=self.guardar(pagina_sintetica())
        pie=[caja(480,258,760,290,'FECHA DE NOTIFICACION:'),caja(480,340,760,372,'HORA DE NOTIFICACION:')]
        m=self.motor(pie,con_vlm=True)
        respuestas=iter(['06/09/2000','15/09/2026'])
        m.read_vlm=lambda im,prompt=None,max_pixels=None: next(respuestas,'ILEGIBLE')
        d,ev=m.leer(ruta)
        self.assertEqual(d['fecha_notificacion'],'2026-09-15'); self.assertNotEqual(ev['confianza'],'CONTRADICTORIA')
        self.assertEqual(ev['fechas_implausibles'],[{'vista':'vlm_color','fechas':['2000-09-06']}])
        # Una fecha anterior a la emisión (01/09/2026) tampoco se propone.
        m=self.motor(pie,con_vlm=True); m.read_vlm=lambda im,prompt=None,max_pixels=None: '01/09/2026'
        m.read_handwritten_lote=lambda imgs: ['nada' for _ in imgs]
        d,ev=m.leer(ruta)
        self.assertIsNone(d['fecha_notificacion']); self.assertTrue(any('2026-09-01' in x['fechas'] for x in ev['fechas_implausibles']))
    def test_zoom_sobre_la_tabla_cuando_no_hay_rotulo(self):
        ruta=self.guardar(pagina_sintetica())
        m=self.motor([])
        base=m.read
        # El detector no ve nada en el pie, pero sí en la pasada con zoom (mitad derecha x2): rótulo en coordenadas del zoom.
        def read(im,digits=False):
            if im.size==(1200,1260): return [caja((480-400)*2,(958-770)*2,(760-400)*2,(990-770)*2,'FECHA DE NOTIFICACION:')]   # zoom desde (400,770)
            return base(im,digits)
        m.read=read
        d,ev=m.leer(ruta)
        self.assertTrue(ev['zoom_rotulos']); self.assertEqual(ev['tipo'],'REC'); self.assertEqual(ev['recuadros'][0]['metodo'],'rotulo_fecha')
        self.assertEqual(d['fecha_notificacion'],'2026-09-15'); self.assertNotIn('recuadro_por_posicion',ev)
    def test_bandas_de_tinta(self):
        im=Image.new('RGB',(400,200),'white'); dr=ImageDraw.Draw(im)
        dr.rectangle((50,30,300,50),fill=(40,60,200)); dr.rectangle((50,120,200,150),fill=(40,60,200)); dr.line((0,100,400,100),fill='black',width=3)
        self.assertEqual(c.bandas_de_tinta(im),[(30,51),(120,151)])
        self.assertEqual(c.bandas_de_tinta(Image.new('RGB',(50,50),'white')),[])
        self.assertEqual(c.zona_cargo_por_posicion(1000,1400),(350,728,780,1176))

class LectorAPITests(unittest.TestCase):
    def lector(self, respuestas):
        l=c.LectorAPI('deepseek',api_key='clave-de-prueba',reintentos=1)
        l.cuerpos=[]
        def peticion(cuerpo):
            l.cuerpos.append(cuerpo); r=respuestas.pop(0)
            if isinstance(r,Exception): raise r
            return {'choices':[{'message':{'content':r}}],'usage':{'total_tokens':420}}
        l._peticion=peticion
        return l
    def test_peticion_y_respuesta(self):
        l=self.lector(['15/09/2026'])
        self.assertEqual(l.read(Image.new('RGB',(60,20),'white')),'15/09/2026')
        cuerpo=l.cuerpos[0]
        self.assertEqual(cuerpo['model'],'deepseek-v4-flash-vision-exp'); self.assertEqual(cuerpo['temperature'],0)
        partes=cuerpo['messages'][0]['content']
        self.assertEqual(partes[0]['type'],'text'); self.assertIn('ILEGIBLE',partes[0]['text'])
        self.assertTrue(partes[1]['image_url']['url'].startswith('data:image/png;base64,'))
        self.assertEqual((l.llamadas,l.errores,l.tokens),(1,0,420))
    def test_reintento_y_error_no_detienen(self):
        l=self.lector([ConnectionError('caída'),'16/09/2026'])
        c.time.sleep=lambda s: None
        self.assertEqual(l.read(Image.new('RGB',(60,20),'white')),'16/09/2026'); self.assertEqual(l.llamadas,2)
        l=self.lector([ConnectionError('caída'),ConnectionError('caída otra vez')])
        self.assertTrue(l.read(Image.new('RGB',(60,20),'white')).startswith('ERROR: ConnectionError')); self.assertEqual(l.errores,1)
    def test_sin_clave_y_proveedor(self):
        import os
        os.environ.pop('DEEPSEEK_API_KEY',None)
        with self.assertRaises(ValueError): c.LectorAPI('deepseek')
        with self.assertRaises(ValueError): c.LectorAPI('otro',api_key='x')
    def test_api_cuenta_como_lector_manuscrito(self):
        v=lambda vista,fechas:{'vista':vista,'fechas':fechas,'texto':''}
        F='2026-09-15'
        self.assertEqual(c.nivel_confianza([v('vlm_color',[F]),v('api_color',[F])],F),('ALTA',['api','vlm']))
        self.assertEqual(c.nivel_confianza([v('color',[F]),v('api_color',[F])],F),('MEDIA',['api','detector']))
        self.assertEqual(c.nivel_confianza([v('api_color',[F]),v('trocr_color',['2026-09-16'])],None)[0],'CONTRADICTORIA')
        class M: vlm=None; api=None
        self.assertIn('no es alcanzable',c.aviso_lectores(M()))
        M.api=type('A',(),{'nombre':'DeepSeek (API)'})(); self.assertIn('DeepSeek (API)',c.aviso_lectores(M())); self.assertIn('ALTA exige',c.aviso_lectores(M()))
    def test_integracion_en_leer(self):
        base=ImagenesSinteticasTests('test_bandas_de_tinta')
        ruta=base.guardar(pagina_sintetica())
        try:
            pie=[caja(480,258,760,290,'FECHA DE NOTIFICACION:'),caja(480,340,760,372,'HORA DE NOTIFICACION:')]
            m=base.motor(pie,con_vlm=True)
            m.api=c.LectorAPI('deepseek',api_key='x'); m.api._peticion=lambda cuerpo: {'choices':[{'message':{'content':'15/09/2026'}}],'usage':{}}
            d,ev=m.leer(ruta)
            vistas={x['vista']:x['fechas'] for x in ev['variantes_fecha']}
            self.assertEqual(vistas['api_color'],['2026-09-15']); self.assertIn('api_contraste',vistas)
            self.assertEqual((d['fecha_notificacion'],ev['confianza']),('2026-09-15','ALTA')); self.assertIn('api',ev['lectores_coinciden'])
            m.api._peticion=lambda cuerpo: (_ for _ in ()).throw(ConnectionError('sin red')); m.api.reintentos=0
            d,ev=m.leer(ruta)
            vistas={x['vista']:x for x in ev['variantes_fecha']}
            self.assertEqual(vistas['api_color']['fechas'],[]); self.assertTrue(vistas['api_color']['texto'].startswith('ERROR'))
            self.assertEqual(d['fecha_notificacion'],'2026-09-15')   # los demás lectores siguen
        finally:
            base.tearDown()

class NotebookTests(unittest.TestCase):
    def setUp(self):
        sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
        import build_notebook
        self.builder=build_notebook
        self.path=Path(__file__).resolve().parents[1]/'notebooks/REC_2026_02_SEPTIEMBRE_Colab.ipynb'
    def test_notebook_has_clean_outputs_and_valid_syntax(self):
        nb=json.loads(self.path.read_text(encoding='utf-8'))
        self.assertEqual(nb['nbformat'],4)
        self.assertEqual(nb['metadata']['accelerator'],'GPU')
        self.assertEqual(nb['metadata']['colab']['gpuType'],'T4')
        cells=[x for x in nb['cells'] if x['cell_type']=='code']
        self.assertEqual(len(cells),7)
        for cell in cells:
            self.assertEqual(cell['outputs'],[]);self.assertIsNone(cell['execution_count'])
            source=''.join(cell['source'])
            ast.parse(source)
            self.assertTrue(source.startswith('# @title '),cell['id'])
        self.assertIn('@param',''.join(cells[0]['source']))
        self.assertIn('def configurar(',''.join(cells[1]['source']))
        self.assertIn('DOBLE_REVISION',''.join(cells[5]['source']))
        self.assertIn('ESCRIBIR_EN_EXCEL_DE_DRIVE',''.join(cells[5]['source']))
        self.assertIn('EXCEL_DRIVE',''.join(cells[1]['source']))
        self.assertIn('LECTOR_VLM',''.join(cells[3]['source']))
        self.assertIn('langchain',''.join(cells[0]['source']))
        self.assertIn('PADDLE_AISLADO',''.join(cells[0]['source']))
        self.assertIn('paddle_python',''.join(cells[3]['source']))
        self.assertIn('medir_acierto',''.join(cells[2]['source']))
        self.assertIn('MUESTRA = 200',''.join(cells[2]['source']))
        self.assertIn('REPROCESAR_MUESTRA',''.join(cells[2]['source']))
        self.assertIn('LECTOR_API',''.join(cells[3]['source'])); self.assertIn('DEEPSEEK_API_KEY',''.join(cells[2]['source']))
    def test_notebook_matches_templates(self):
        nb=json.loads(self.path.read_text(encoding='utf-8'))
        self.assertEqual(nb,self.builder.render(),'Ejecuta python scripts/build_notebook.py y confirma el cuaderno.')
if __name__=='__main__':unittest.main()
