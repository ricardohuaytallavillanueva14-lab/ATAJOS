from pathlib import Path
from datetime import datetime, date
from collections import defaultdict, Counter
from zoneinfo import ZoneInfo
import base64, copy, csv, hashlib, html, io, json, os, re, time, unicodedata
import zipfile, posixpath, uuid
from PIL import Image, ImageOps
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import to_excel
from lxml import etree as ET

VERSION = '2026-10-23-v2'
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
N = {'s': NS}
EXTENSIONES = {'.jpg', '.jpeg', '.png', '.webp', '.tif', '.tiff', '.bmp'}

# Alcance fijo del lote. Se cambia solo con configurar(); lo usan el índice, la
# evaluación y la escritura del XLSX, de modo que las tres barreras coinciden.
HOJA = '2026'
FECHA_EMISION = date(2026, 9, 2)
# Días tras la emisión dentro de los cuales una notificación se considera habitual.
# Fuera de la ventana no se bloquea: se marca una alerta para verificar año y mes.
VENTANA_DIAS = 60
MODELOS_VLM = {'qwen': 'Qwen/Qwen2.5-VL-3B-Instruct'}

try:
    from tqdm.auto import tqdm
except ImportError:  # pragma: no cover - tqdm viene instalado en Colab
    tqdm = None


def configurar(hoja=None, fecha_emision=None, ventana_dias=None):
    """Fija la hoja anual y la fecha de emisión REC que delimitan el lote.

    Devuelve (hoja, fecha_emision). Una carpeta de resultados queda ligada a
    esta configuración mediante proyecto.json. ventana_dias ajusta la alerta
    de plausibilidad (no bloquea).
    """
    global HOJA, FECHA_EMISION, VENTANA_DIAS
    if ventana_dias is not None:
        if int(ventana_dias) < 1:
            raise ValueError('La ventana de plausibilidad debe ser de al menos un día.')
        VENTANA_DIAS = int(ventana_dias)
    nueva_hoja = HOJA if hoja is None else str(hoja).strip()
    if not re.fullmatch(r'20\d{2}', nueva_hoja):
        raise ValueError('La hoja debe ser un año de cuatro cifras, por ejemplo 2026.')
    nueva_fecha = FECHA_EMISION if fecha_emision is None else fecha(fecha_emision)
    if nueva_fecha is None:
        raise ValueError('Fecha de emisión inválida. Usa el formato DD/MM/AAAA.')
    if str(nueva_fecha.year) != nueva_hoja:
        raise ValueError(f'La fecha de emisión {nueva_fecha:%d/%m/%Y} no pertenece a la hoja {nueva_hoja}.')
    HOJA, FECHA_EMISION = nueva_hoja, nueva_fecha
    return HOJA, FECHA_EMISION


def vista_tinta(im, min_fraccion=0.002):
    """Aísla los trazos de bolígrafo azul (manuscritos) del texto y las líneas impresas en negro.

    Devuelve (imagen RGB con solo la tinta sobre blanco, fracción de píxeles de tinta) o
    (None, fracción) si casi no hay tinta azul: entonces el manuscrito es negro o no existe.
    """
    import numpy as np
    hsv = np.asarray(im.convert('RGB').convert('HSV'), dtype=np.int16)
    h, sat, val = hsv[..., 0], hsv[..., 1], hsv[..., 2]
    # Azul/violeta de bolígrafo en la escala 0-255 de PIL: aprox. 145-200 (el cian del escáner
    # queda en ~128 y se excluye). Impresión negra: saturación baja.
    mascara = (sat > 45) & (val > 25) & (val < 245) & (h >= 145) & (h <= 200)
    fraccion = float(mascara.mean())
    if fraccion < min_fraccion:
        return None, fraccion
    salida = np.full(hsv.shape, 255, dtype=np.uint8)
    salida[mascara] = 0
    return Image.fromarray(salida, 'RGB'), fraccion


# Lectores entrenados en escritura manual. El detector impreso (PaddleOCR/EasyOCR) lee mal
# los manuscritos, así que su coincidencia no basta para la confianza ALTA.
LECTORES_MANUSCRITOS = ('trocr', 'vlm', 'deepseek')


def familia_lector(vista):
    """Agrupa las vistas por lector independiente: detector (impreso), trocr, vlm, deepseek, acta."""
    for prefijo in LECTORES_MANUSCRITOS + ('acta',):
        if str(vista).startswith(prefijo):
            return prefijo
    return 'detector'


def nivel_confianza(variantes, propuesta):
    """Clasifica la propuesta de fecha según qué lectores independientes coinciden.

    ALTA: al menos dos lectores manuscritos distintos (TrOCR, VLM) leyeron exactamente la
        misma fecha y ningún lector leyó otra. El detector impreso no cuenta para ALTA:
        detector + un manuscrito = MEDIA. Sin el lector VLM, ALTA no es alcanzable.
    MEDIA: hay propuesta, nadie la contradice, pero no llegan a dos lectores manuscritos.
    CONTRADICTORIA: hay lecturas de fechas distintas. SIN_FECHA: nadie leyó una fecha.
    Las lecturas marcadas como corregidas no cuentan para ningún nivel.
    La fecha del acta ('acta') no interviene: solo genera alerta (ver alertas_cruce).
    Ningún nivel sustituye el cotejo visual: solo ordena y documenta la revisión.
    """
    por_familia = defaultdict(set)
    for v in variantes:
        familia = familia_lector(v.get('vista'))
        if familia == 'acta':
            continue
        por_familia[familia].update(v.get('fechas') or [])
    distintas = {d for ds in por_familia.values() for d in ds}
    if len(distintas) > 1:
        return 'CONTRADICTORIA', sorted(f for f, ds in por_familia.items() if ds == {propuesta})
    if propuesta is None:
        return 'SIN_FECHA', []
    coinciden = sorted(f for f, ds in por_familia.items() if ds == {propuesta})
    manuscritos = [f for f in coinciden if f in LECTORES_MANUSCRITOS]
    return ('ALTA' if len(manuscritos) >= 2 else 'MEDIA'), coinciden


def aviso_lectores(motor):
    """Texto para pantalla sobre qué lectores manuscritos hay y si ALTA es alcanzable."""
    manuscritos = ['TrOCR'] + (['Qwen2.5-VL'] if getattr(motor, 'vlm', None) is not None else [])
    if len(manuscritos) < 2:
        error = getattr(motor, 'vlm_error', None)
        return ((f'El lector VLM no se cargó ({error}). ' if error else '') + 'Lectores manuscritos activos: ' + ', '.join(manuscritos) + '. Con un solo lector manuscrito la '
                'confianza ALTA no es alcanzable (es intencional): todas las fechas quedarán como máximo en MEDIA '
                'y requieren el mismo cotejo visual. Activa LECTOR_VLM para poder alcanzar ALTA.')
    return 'Lectores manuscritos activos: ' + ', '.join(manuscritos) + '. ALTA exige que ambos coincidan.'


def alertas(data, target=None, evidencia=None):
    """Avisos de plausibilidad que no bloquean la aprobación pero exigen atención.

    evidencia: evidencia_ocr del registro; si trae 'fecha_acta' (la fecha escrita en
    'del día __/__/__' del acta) distinta de la propuesta, se añade una alerta. El cruce
    no cambia la confianza: en DOMICILIO CERRADO el acta lleva la fecha del aviso previo
    (día anterior), así que una diferencia puede ser legítima; en NEGATIVA debería coincidir.
    """
    avisos = []
    d = fecha((data or {}).get('fecha_notificacion'))
    if d is None:
        return avisos
    acta = [a for a in ((evidencia or {}).get('fecha_acta') or []) if fecha(a)]
    if acta and d.isoformat() not in acta:
        avisos.append('La fecha escrita en el acta («del día …») es ' + ', '.join(f'{fecha(a):%d/%m/%Y}' for a in acta)
                      + f' y difiere de la del recuadro ({d:%d/%m/%Y}). En DOMICILIO CERRADO suele ser el aviso previo del día anterior; en NEGATIVA debe coincidir. Verifica en la foto.')
    em = fecha(target['emision']) if target else fecha((data or {}).get('fecha_emision'))
    if em and (d - em).days > VENTANA_DIAS:
        avisos.append(f'La notificación es {(d - em).days} días posterior a la emisión (ventana habitual: {VENTANA_DIAS}). Verifica mes y año en la foto.')
    if d.weekday() == 6:
        avisos.append('La fecha cae en domingo; confirma el día en la foto.')
    if (data or {}).get('resolucion') in (None, ''):
        avisos.append('No se leyó el número de resolución; confirma en la foto que es la REC UNO.')
    if (data or {}).get('fecha_emision') in (None, ''):
        avisos.append('No se leyó la fecha de emisión impresa; confirma en la foto que la REC es del lote.')
    nombre_doc = (data or {}).get('nombre'); nombre_fila = (target or {}).get('nombre')
    if nombre_doc and nombre_fila and similitud_nombres(nombre_doc, nombre_fila) < 0.5:
        avisos.append(f'El nombre leído ({nombre_doc[:60]}) no se parece al de la fila ({str(nombre_fila)[:60]}). Confirma que la foto corresponde a esa fila.')
    return avisos


def similitud_nombres(a, b):
    """Similitud entre nombres (0-1) tolerante al orden y a apellidos compuestos: comparación de palabras."""
    import difflib
    pa = sorted(set(re.findall(r'[a-z]{2,}', norm_texto(a)))); pb = sorted(set(re.findall(r'[a-z]{2,}', norm_texto(b))))
    if not pa or not pb:
        return 0.0
    comunes = sum(1 for w in pa if any(difflib.SequenceMatcher(None, w, v).ratio() >= 0.85 for v in pb))
    return comunes / min(len(pa), len(pb))


def norm_texto(value):
    """Minúsculas sin tildes, conservando espacios (para comparar nombres por palabras)."""
    s = unicodedata.normalize('NFKD', str(value or ''))
    return re.sub(r'[^a-z0-9 ]', ' ', ''.join(c for c in s if not unicodedata.combining(c)).lower())


def _progreso(iterable, total, desc):
    """Barra de progreso con tiempo estimado; si no hay tqdm, imprime cada 25."""
    if tqdm is not None:
        return tqdm(iterable, total=total, desc=desc, unit='foto', dynamic_ncols=True)
    def plano():
        for i, item in enumerate(iterable, 1):
            yield item
            if i % 25 == 0 or i == total:
                print(f'{desc}: {i}/{total}', flush=True)
    return plano()

def ahora():
    return datetime.now(ZoneInfo('America/Lima')).isoformat(timespec='seconds')

def sha256(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()

def atomic_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    with tmp.open('w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)

def norm(value):
    s = unicodedata.normalize('NFKD', str(value or ''))
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in s if not unicodedata.combining(c)).lower())

def codigo(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value or '').strip()
    return str(int(s)) if re.fullmatch(r'\d+', s) else None

def expediente(value):
    s = str(value or '').strip().upper().replace('–', '-').replace('—', '-')
    s = re.sub(r'\s+', '', s)
    m = re.fullmatch(r'(\d{2})-(\d{3})-(\d{1,12})(?:-(20\d{2})(?:-SAT-H/DCC(?:-[A-Z]+)?)?)?', s)
    if not m:
        return None
    return f'{m[1]}-{m[2]}-{int(m[3])}', int(m[4]) if m[4] else None

def fecha(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value or '').strip()
    for fmt in ('%Y-%m-%dT%H:%M:%S', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d/%m/%Y %H:%M:%S'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def texto_celda(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value

def build_index(excel, hojas):
    if tuple(hojas) != (HOJA,):
        raise ValueError(f'Solo se permite la hoja {HOJA}.')
    if not zipfile.is_zipfile(excel):
        raise ValueError('El archivo no es un XLSX válido o está incompleto. Ábrelo en Excel y guarda una copia .xlsx antes de continuar.')
    wb = load_workbook(excel, read_only=True, data_only=False, keep_links=True)
    index, resumen = defaultdict(list), []
    try:
        for sheet_name in hojas:
            if sheet_name not in wb.sheetnames or not re.fullmatch(r'20\d{2}', sheet_name):
                raise ValueError(f'Hoja anual inexistente o no válida: {sheet_name}')
            ws = wb[sheet_name]
            # Leer todas las filas, incluso las que quedan fuera del autofiltro.
            rows = ws.iter_rows(values_only=True)
            headers = [norm(v) for v in next(rows)]
            def unique(pred, label):
                found = [i for i, h in enumerate(headers) if pred(h)]
                if len(found) != 1:
                    raise ValueError(f'{sheet_name}: no se identificó una sola columna {label}. Encabezados: {headers}')
                return found[0]
            exp_col = unique(lambda h: h in {'expdiente','expediente','nexpediente','ndeexpediente','nrodeexpediente','ndeexp'}, 'expediente')
            date_col = unique(lambda h: h in {'ffechadenotificacionrec','fechadenotificacionrec','fechanotificacionrec','fechanotifrec','fechadenotificacion'}, 'notificación REC')
            code_col = unique(lambda h: h in {'codsigeti','codigodeobligsigeti'}, 'código SIGETI')
            name_cols = [i for i,h in enumerate(headers) if h in {'nombre','obligado'}]
            emission_cols = [i for i,h in enumerate(headers) if h in {'fechaemisionrec','fechadeemisionderec','fechaemisrec'}]
            if date_col != 5 or emission_cols != [4]:
                raise ValueError('Se esperaba emisión REC en E y notificación REC en F.')
            count = 0
            for row_num, row in enumerate(rows, 2):
                def at(i): return row[i] if i < len(row) else None
                parsed = expediente(at(exp_col))
                if not parsed:
                    continue
                key, explicit_year = parsed
                year = explicit_year or int(sheet_name)
                if year != int(sheet_name):
                    continue
                item = dict(hoja=sheet_name, fila=row_num, celda=f'{get_column_letter(date_col+1)}{row_num}',
                            expediente=str(at(exp_col)), codigo=codigo(at(code_col)),
                            nombre=str(at(name_cols[0]) or '') if name_cols else '',
                            fecha_actual=texto_celda(at(date_col)),
                            emision=texto_celda(at(emission_cols[0])) if emission_cols else None)
                index[(key,year)].append(item)
                count += 1
            resumen.append({'hoja':sheet_name,'columna_fecha':get_column_letter(date_col+1),'expedientes':count})
        epoch = wb.epoch
    finally:
        wb.close()
    return dict(index), resumen, epoch

def normalizar_lectura(data):
    if not isinstance(data, dict):
        raise ValueError('La respuesta no es un objeto JSON.')
    fields = ['expediente','codigo','resolucion','nombre','fecha_notificacion','fecha_emision','fecha_texto','observaciones']
    out = {}
    for field in fields:
        v = data.get(field)
        if v is not None and not isinstance(v, str):
            raise ValueError(f'Campo {field} debe ser texto o null.')
        out[field] = v.strip() if isinstance(v, str) else None
    return out

def evaluar(data, index):
    issues, target = [], None
    parsed = expediente(data.get('expediente'))
    if not parsed or parsed[1] is None:
        issues.append('Expediente completo y año ilegibles o incompletos.')
    else:
        matches = index.get(parsed, [])
        if len(matches) != 1:
            issues.append(f'Coincidencias de expediente/año en Excel: {len(matches)}. Debe existir exactamente una.')
        else:
            target = matches[0]
            if target['hoja'] != HOJA or fecha(target['emision']) != FECHA_EMISION:
                issues.append(f'Fuera del alcance: solo {HOJA} con E = {FECHA_EMISION:%d/%m/%Y}.')
            if target['fecha_actual'] not in (None, ''):
                issues.append('La notificación ya está registrada. Se conserva sin modificar.')
            if not codigo(data.get('codigo')) or codigo(data.get('codigo')) != target['codigo']:
                issues.append('El código del contribuyente no coincide con el Excel o no se puede leer.')
            em_doc = fecha(data.get('fecha_emision'))
            if em_doc and fecha(target['emision']) and em_doc != fecha(target['emision']):
                issues.append(f"La fecha de emisión impresa en el documento ({em_doc:%d/%m/%Y}) no coincide con la columna E de la fila ({fecha(target['emision']):%d/%m/%Y}).")
    if data.get('resolucion') not in (None, '') and norm(data.get('resolucion')) not in {'uno','1','01','resolucionnuno','resolucionuno'}:
        issues.append('Esta versión solo llena la notificación de la REC UNO. Otra resolución requiere revisión aparte.')
    d = fecha(data.get('fecha_notificacion'))
    if d is None:
        issues.append('Fecha de notificación incompleta, ilegible o inválida.')
    else:
        if d > datetime.now(ZoneInfo('America/Lima')).date():
            issues.append('La fecha de notificación está en el futuro.')
        em = fecha(data.get('fecha_emision'))
        excel_em = fecha(target['emision']) if target else None
        if (em and d < em) or (excel_em and d < excel_em):
            issues.append('La notificación es anterior a la emisión de la REC.')
        if target and target['fecha_actual'] not in (None, '') and fecha(target['fecha_actual']) != d:
            issues.append('La celda ya contiene otra fecha o una fórmula. No se sobrescribe.')
    return target, issues

def parse_header(text):
    # No sustituir O por 0 ni completar dígitos: un identificador dudoso se revisa.
    clean = text.upper().replace('–','-').replace('—','-')
    pattern = r'\b(\d{2}\s*-\s*\d{3}\s*-\s*\d{1,12}\s*-\s*20\d{2})(?:\s*-\s*SAT\s*-\s*H\s*/\s*DCC(?:\s*-\s*[A-Z]+)?)?'
    found = {re.sub(r'\s+','',m.group(0)) for m in re.finditer(pattern,clean)}
    exp = next(iter(found)) if len(found)==1 else None
    codes = re.findall(r'(?:CONTRIBUYENTE|SIGETI)\s*[:.;°ºN\s]*([0-9]{3,15})\b',clean)
    codes = {codigo(c) for c in codes}
    code = next(iter(codes)) if len(codes)==1 else None
    resolution = re.findall(r'RESOLUCI[OÓ0Ò]N\s*(?:N[°ºO.]?\s*)?(UNO|DOS|TRES|\d+)\b',clean)
    if not resolution:
        # Lectura aproximada: una palabra parecida a RESOLUCION seguida de UNO/DOS/TRES o un número.
        tokens = re.findall(r'[A-ZÁÉÍÓÚÑ0-9°º]+', clean)
        for i, tok in enumerate(tokens[:-1]):
            if parecido(tok, 'resolucion', umbral=0.7, prefijo=5):
                for sig in tokens[i+1:i+3]:
                    if sig in ('UNO','DOS','TRES') or sig.isdigit():
                        resolution.append(sig); break
    return exp,code,resolution[0] if len(set(resolution))==1 else None

MESES = ['enero','febrero','marzo','abril','mayo','junio','julio','agosto','setiembre','octubre','noviembre','diciembre']


def fecha_en_letras(text):
    """'Ayacucho, 02 de setiembre de 2026' -> '2026-09-02'. Tolera 'septiembre' y meses mal leídos (similitud >= 0,75)."""
    import difflib
    for d, mes, y in re.findall(r'(?<!\d)(\d{1,2})\s+DE\s+([A-ZÁÉÍÓÚÑ]{3,12})\s+DE[LE]?\s+(20\d{2})', str(text or '').upper()):
        m = norm(mes).replace('septiembre', 'setiembre')
        candidatos = [(difflib.SequenceMatcher(None, m, nombre).ratio(), i + 1) for i, nombre in enumerate(MESES)]
        ratio, numero = max(candidatos)
        if ratio >= 0.75:
            f = _fecha_iso(d, numero, y)
            if f:
                return f
    return None


def parse_header_extra(text):
    """Nombre del obligado y fecha de emisión impresa del encabezado de la REC.

    El nombre sale de la línea 'OBLIGADO (A) : ...'; la fecha, de 'Ayacucho, 02 de setiembre de 2026'
    o de una fecha numérica en una línea que mencione emisión. Devuelve (nombre, fecha_iso).
    """
    nombre = None
    emision = None
    for linea in str(text or '').splitlines():
        u = linea.upper()
        m = re.search(r'OBL[I1L]GAD[O0]\s*\(?\s*A?\s*\)?\s*[:.;]?\s*(.+)$', u)
        if m and nombre is None:
            valor = re.sub(r'\s+', ' ', m.group(1)).strip(' :.;-')
            valor = re.split(r'\s+(?:DOMICILIO|MATERIA|AUXILIAR|C[ÓO]DIGO)\b', valor)[0].strip()
            if len(valor) >= 5 and re.search(r'[A-ZÁÉÍÓÚÑ]{2,}', valor):
                nombre = valor
        if emision is None:
            emision = fecha_en_letras(u)
            if emision is None and 'EMISI' in norm(u).upper():
                ds = dates_in_text(u)
                emision = ds[0] if len(ds) == 1 else None
    return nombre, emision


def _fecha_iso(day, month, year):
    try:
        return date(int(year), int(month), int(day)).isoformat()
    except ValueError:
        return None


def _fechas_exactas(text):
    """Patrones de fecha sin ninguna corrección de caracteres."""
    matches = re.findall(r'(?<!\d)(\d{1,2})\s*[/ .-]\s*(\d{1,2})\s*[/ .-]\s*(20\d{2})(?!\d)', text)
    # Con la tinta aislada las barras finas pueden desaparecer: 8 cifras seguidas DDMMAAAA.
    matches += re.findall(r'(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)', text)
    # Año de dos cifras (16/09/26) solo con separadores, para no confundirlo con otros números.
    matches += [(d, m, '20' + y) for d, m, y in re.findall(r'(?<!\d)(\d{1,2})\s*[/.-]\s*(\d{1,2})\s*[/.-]\s*(\d{2})(?!\d)', text)]
    matches += [(d, m, '20' + y) for d, m, y in re.findall(r'(?<!\d)(\d{1,2}) +(\d{1,2}) +(\d{2})(?!\d)', text)]
    return {f for f in (_fecha_iso(*m) for m in matches) if f}


_CONFUSIONES = str.maketrans({'O': '0', 'o': '0', 'I': '1', 'l': '1', '|': '1'})


def extraer_fechas(text):
    """Fechas del texto en dos listas.

    exactas: patrones DD/MM/AAAA (o DD/MM/AA con separadores) tal cual, incluso con espacios
        entre dígitos ('1 5 / 0 9 / 2 0 2 6'), que se eliminan porque no cambian ninguna cifra.
    corregidas: candidatas obtenidas sustituyendo O→0 e I/l→1 solo dentro de un patrón de
        fecha, o recortando un año de cinco cifras que contiene '20AA'. Son ambiguas: se
        muestran al revisor, pueden proponerse si no hay nada mejor y nunca cuentan para la
        confianza ALTA. Nunca se completan dígitos que faltan ni se adivina el año.
    """
    text = str(text or '')
    compacto = re.sub(r'(?<=\d) +(?=\d)', '', text)
    exactas = _fechas_exactas(text) | _fechas_exactas(compacto)
    corregidas = {}
    patron = r'(?<![0-9A-Za-z])([0-9OoIl|]{1,2})\s*[/.-]\s*([0-9OoIl|]{1,2})\s*[/.-]\s*([0-9OoIl|]{4}|[0-9OoIl|]{2})(?![0-9A-Za-z])'
    for fuente in (text, compacto):
        for m in re.finditer(patron, fuente):
            bruto = m.group(0)
            if sum(ch in 'OoIl|' for ch in bruto) == 0 or sum(ch in 'OoIl|' for ch in bruto) > 2:
                continue
            d, mo, y = (g.translate(_CONFUSIONES) for g in m.groups())
            if not (d.isdigit() and mo.isdigit() and y.isdigit()):
                continue
            if len(y) == 2:
                y = '20' + y
            elif not y.startswith('20'):
                continue
            f = _fecha_iso(d, mo, y)
            if f and f not in exactas:
                corregidas.setdefault(f, {'fecha': f, 'regla': 'O→0, I/l→1', 'texto': bruto.strip()})
        for m in re.finditer(r'(?<!\d)(\d{1,2})\s*[/.-]\s*(\d{1,2})\s*[/.-]\s*(\d{5})(?!\d)', fuente):
            d, mo, y5 = m.groups()
            for y in (y5[:4], y5[1:]):
                if y.startswith('20'):
                    f = _fecha_iso(d, mo, y)
                    if f and f not in exactas:
                        corregidas.setdefault(f, {'fecha': f, 'regla': 'año de 5 cifras', 'texto': m.group(0).strip()})
                    break
    return {'exactas': sorted(exactas), 'corregidas': [corregidas[k] for k in sorted(corregidas)]}


def dates_in_text(text):
    """Fechas exactas (sin correcciones) del texto; véase extraer_fechas()."""
    return extraer_fechas(text)['exactas']

def seleccionar_muestra(paths, raiz, cantidad=30, semilla=2026):
    """Muestra reproducible repartida por carpeta principal y subcarpeta final."""
    import random
    from collections import deque
    if cantidad < 1:
        raise ValueError('La muestra debe contener al menos una foto.')
    raiz=Path(raiz).resolve()
    grupos=defaultdict(lambda:defaultdict(list))
    for p in sorted(set(Path(p).resolve() for p in paths)):
        rel=p.relative_to(raiz)
        principal=rel.parts[0] if len(rel.parts)>1 else '(raiz)'
        grupos[principal][rel.parent.as_posix()].append(p)
    rng=random.Random(semilla)
    principales=list(sorted(grupos));rng.shuffle(principales)
    colas={}
    for principal in principales:
        hojas=list(sorted(grupos[principal]));rng.shuffle(hojas)
        filas=[]
        for hoja in hojas:
            fotos=grupos[principal][hoja];rng.shuffle(fotos)
            filas.append(deque(fotos))
        colas[principal]=deque(filas)
    turnos=deque(principales); seleccion=[]
    while turnos and len(seleccion)<cantidad:
        principal=turnos.popleft()
        fila=colas[principal].popleft()
        seleccion.append(fila.popleft())
        if fila: colas[principal].append(fila)
        if colas[principal]: turnos.append(principal)
    return seleccion


def parecido(texto, objetivo, umbral=0.65, prefijo=4):
    """Coincidencia aproximada tolerante al OCR: mismo inicio y similitud global mínima."""
    import difflib
    t = norm(texto)
    if not t or t[:prefijo] != objetivo[:prefijo]:
        return 0.0
    r = difflib.SequenceMatcher(None, t, objetivo).ratio()
    return r if r >= umbral else 0.0


def parecido_rotulo(texto, inicio, otro):
    """Puntuación de un rótulo 'FECHA DE NOTIFICACIÓN' u 'HORA DE NOTIFICACIÓN' leído con errores.

    inicio: primera palabra esperada ('fecha' u 'hora'); otro: la del rótulo contrario.
    Acepta: inicio reconocible (mismas 4 letras o similitud >= 0,6) con similitud global
    >= 0,65 y una cola parecida a NOTIFICACION (>= 0,55); o, aunque falle el inicio, una
    cola muy parecida a NOTIFICACION (>= 0,65), con puntuación menor por ser ambigua.
    Rechaza lo que empieza como el rótulo contrario y los rótulos de otras fechas
    (emisión, vencimiento), cuya cola no se parece a NOTIFICACION.
    """
    import difflib
    t = norm(texto)
    if not t:
        return 0.0
    R = lambda a, b: difflib.SequenceMatcher(None, a, b).ratio()
    inicio_ok = t[:4] == inicio[:4] or R(t[:len(inicio)], inicio) >= 0.6
    otro_ok = t[:4] == otro[:4] or R(t[:len(otro)], otro) >= 0.75
    if otro_ok and not t[:4] == inicio[:4]:
        return 0.0
    cola = max([R(t[-n:], 'notificacion') for n in range(10, 15) if len(t) >= n] or [R(t, 'notificacion')])
    global_r = R(t, inicio + 'denotificacion')
    if inicio_ok and global_r >= 0.65 and cola >= 0.55:
        return global_r
    if not inicio_ok and cola >= 0.65:
        return 0.5 + cola / 10  # ambiguo: solo la palabra NOTIFICACION
    return 0.0


def agrupar_lineas(rows, tolerancia=0.6):
    """Une detecciones que están en la misma línea (p. ej. 'FECHA DE' + 'NOTIFICACIÓN:')."""
    items = []
    for r in rows:
        xs = [x for x, y in r['box']]; ys = [y for x, y in r['box']]
        items.append({'x1': min(xs), 'x2': max(xs), 'y1': min(ys), 'y2': max(ys), 'text': r['text'], 'score': r.get('score', 0)})
    items.sort(key=lambda i: (i['y1'], i['x1']))
    lineas = []
    for it in items:
        h = max(1, it['y2'] - it['y1']); cy = (it['y1'] + it['y2']) / 2
        for ln in lineas:
            lcy = (ln['y1'] + ln['y2']) / 2; lh = max(1, ln['y2'] - ln['y1'])
            if abs(cy - lcy) <= tolerancia * min(h, lh) and (it['x1'] - ln['x2']) < 3 * lh:
                ln['text'] += ' ' + it['text']
                ln['x1'] = min(ln['x1'], it['x1']); ln['x2'] = max(ln['x2'], it['x2'])
                ln['y1'] = min(ln['y1'], it['y1']); ln['y2'] = max(ln['y2'], it['y2'])
                break
        else:
            lineas.append(dict(it))
    return lineas


def localizar_recuadro_fecha(rows, ancho, alto):
    """Ubica la casilla manuscrita bajo 'FECHA DE NOTIFICACIÓN' aunque el rótulo se lea mal.

    rows: detecciones del pie de página con coordenadas en el mismo sistema que ancho/alto.
    Devuelve (roi (x1,y1,x2,y2), método) o (None, motivo). Estrategias, en orden:
    1. Rótulo FECHA (aproximado) y rótulo HORA debajo: casilla entre ambos.
    2. Solo rótulo FECHA: casilla de ~2,6 alturas de línea bajo él.
    3. Solo rótulo HORA: casilla justo encima, de la misma altura que la de HORA.
    """
    lineas = agrupar_lineas(rows)
    fechas = sorted(((parecido_rotulo(l['text'], 'fecha', 'hora'), l) for l in lineas), key=lambda t: -t[0])
    horas = sorted(((parecido_rotulo(l['text'], 'hora', 'fecha'), l) for l in lineas), key=lambda t: -t[0])
    fecha = fechas[0][1] if fechas and fechas[0][0] > 0 else None
    fecha_ambigua = fecha is not None and fechas[0][0] < 0.65
    hora = None
    for score, l in horas:
        if score <= 0:
            break
        if l is fecha:
            continue
        if fecha is None or (l['y1'] > fecha['y2'] and abs(l['x1'] - fecha['x1']) < 0.25 * ancho):
            hora = l; break
    if fecha_ambigua and hora is None:
        # Solo se leyó 'NOTIFICACION' suelto: puede ser el rótulo de FECHA o el de HORA.
        return None, 'Rótulo ambiguo (solo se reconoció NOTIFICACIÓN, sin FECHA/HORA ni rótulo HORA debajo).'
    if fecha is None and hora is None:
        return None, 'No se localizó ningún rótulo FECHA/HORA DE NOTIFICACIÓN.'
    if fecha is not None:
        lh = max(10, fecha['y2'] - fecha['y1'])
        top = fecha['y2']
        bottom = hora['y1'] if hora is not None and top < hora['y1'] < top + lh * 6 else min(alto, top + lh * 2.6)
        x1, x2 = fecha['x1'], fecha['x2']
        metodo = 'rotulos_fecha_y_hora' if hora is not None else 'rotulo_fecha'
    else:
        lh = max(10, hora['y2'] - hora['y1'])
        bottom = hora['y1']; top = max(0, bottom - lh * 2.6)
        x1, x2 = hora['x1'], hora['x2']
        metodo = 'rotulo_hora'
    roi = (max(0, int(x1 - 0.015 * ancho)), max(0, int(top)), min(ancho, int(x2 + 0.02 * ancho)), min(alto, int(bottom)))
    if roi[3] - roi[1] < 8 or roi[2] - roi[0] < 20:
        return None, 'El recorte de fecha es demasiado pequeño.'
    return roi, metodo


def fecha_acta_negativa(rows):
    """Fecha manuscrita del 'Acta de negativa' ('Siendo las __ horas, del día __/__/__'): control cruzado."""
    out = set()
    for l in agrupar_lineas(rows):
        t = norm(l['text'])
        if 'deldia' in t or 'siendolas' in t:
            out.update(dates_in_text(l['text']))
    return sorted(out)


def clasificar_documento(head_text, foot_rows, ancho, alto, exp):
    """REC completa, REC sin cargo (media página), cargo sin encabezado o anexo (AVISO PREVIO / CEDULÓN)."""
    cabecera = norm(head_text)
    es_rec = exp is not None or 'expedientecoactivo' in cabecera or 'ejecucioncoactiva' in cabecera
    anexo = ('avisoprevio' in cabecera or 'cedulon' in cabecera) and not es_rec
    roi, _ = localizar_recuadro_fecha(foot_rows, ancho, alto)
    if anexo:
        return 'ANEXO'
    if es_rec and roi:
        return 'REC'
    if es_rec:
        return 'REC_SIN_CARGO'
    if roi:
        return 'CARGO_SIN_ENCABEZADO'
    return 'DESCONOCIDO'


def probar_paddle(timeout=600):
    """Comprueba en un subproceso que PaddleOCR puede inferir; una caída (segfault) no mata el cuaderno."""
    import subprocess, sys
    codigo_prueba = (
        "import numpy as np\n"
        "from paddleocr import PaddleOCR\n"
        "o=PaddleOCR(lang='es',ocr_version='PP-OCRv5',device='cpu',use_doc_orientation_classify=False,"
        "use_doc_unwarping=False,use_textline_orientation=False,enable_mkldnn=False)\n"
        "im=np.full((96,320,3),255,np.uint8); im[30:60,20:300]=0\n"
        "list(o.predict(im)); print('PADDLE_OK')\n")
    try:
        r = subprocess.run([sys.executable, '-c', codigo_prueba], capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        return False, 'tiempo agotado'
    if r.returncode == 0 and 'PADDLE_OK' in r.stdout:
        return True, 'ok'
    motivo = (r.stderr or r.stdout).strip().splitlines()[-1:] or [f'código de salida {r.returncode}']
    return False, f'código {r.returncode}: {motivo[0][:200]}'


class MotorOCR:
    """Detector de texto impreso (PaddleOCR o EasyOCR) + TrOCR manuscrito + VLM opcional."""
    def __init__(self, motor='auto', gpu=False, trocr_model='microsoft/trocr-base-handwritten',
                 paddle_gpu=False, modo_rapido=False, cpu_threads=None, vlm=None, orientacion_auto=False):
        """Carga los modelos una sola vez.

        motor: 'auto' prueba PaddleOCR en un subproceso y, si no puede inferir
            (p. ej. caída en Python 3.13), usa EasyOCR; 'paddle' o 'easyocr' lo fijan.
            Se aceptan los nombres antiguos 'paddle_trocr' y 'easyocr'.
        gpu: usa CUDA para EasyOCR, TrOCR y el VLM si está disponible.
        paddle_gpu: intenta ejecutar PaddleOCR en GPU (requiere paddlepaddle-gpu);
            si falla, vuelve a CPU y lo informa.
        modo_rapido: detector PP-OCRv5 mobile, varias veces más rápido en CPU.
        cpu_threads: hilos para Paddle en CPU; por defecto todos los del entorno.
        vlm: tercer lector independiente de la fecha manuscrita: 'qwen'
            (Qwen2.5-VL-3B-Instruct, ~7 GB en GPU T4) o None.
        orientacion_auto: PaddleOCR corrige fotos giradas 90/180/270 grados.
        """
        motor = {'paddle_trocr': 'paddle'}.get(motor, motor)
        if motor not in ('auto', 'paddle', 'easyocr'):
            raise ValueError('Motor desconocido: ' + str(motor) + ". Usa 'auto', 'paddle' o 'easyocr'.")
        if vlm not in (None, '', *MODELOS_VLM):
            raise ValueError('VLM desconocido: ' + str(vlm) + '. Opciones: ' + ', '.join(MODELOS_VLM))
        import torch
        self.paddle = None; self.reader = None; self.trocr = None; self.vlm = None
        self.opciones = dict(motor=motor, gpu=gpu, trocr_model=trocr_model, paddle_gpu=paddle_gpu,
                             modo_rapido=modo_rapido, cpu_threads=cpu_threads, vlm=vlm or None,
                             orientacion_auto=orientacion_auto)
        self.gpu = bool(gpu and torch.cuda.is_available())
        detector = motor
        if motor == 'auto':
            ok, motivo = probar_paddle()
            detector = 'paddle' if ok else 'easyocr'
            if not ok:
                print('PaddleOCR no puede inferir en este entorno (' + motivo + '); se usará EasyOCR.')
        if detector == 'paddle':
            self._cargar_paddle(paddle_gpu, modo_rapido, cpu_threads, orientacion_auto)
        else:
            import easyocr
            self.reader = easyocr.Reader(['es', 'en'], gpu=self.gpu)
            self.detector_nombre = 'EasyOCR (' + ('gpu' if self.gpu else 'cpu') + ')'
        from transformers import TrOCRProcessor, VisionEncoderDecoderModel
        self.trocr_device = 'cuda' if self.gpu else 'cpu'
        self.processor = TrOCRProcessor.from_pretrained(trocr_model)
        self.trocr = VisionEncoderDecoderModel.from_pretrained(trocr_model).to(self.trocr_device).eval()
        self.nombre = self.detector_nombre + ' + ' + trocr_model + ' (' + self.trocr_device + ')'
        self.vlm_error = None
        if vlm:
            self._activar_vlm(MODELOS_VLM[vlm])

    def _activar_vlm(self, model_name):
        """Carga el VLM; si falla (memoria, pesos, versión), sigue sin él y lo dice con claridad."""
        try:
            self._cargar_vlm(model_name, self.gpu)
            self.nombre += ' + ' + model_name + ' (' + self.vlm_device + ')'
        except Exception as e:
            self.vlm = None
            self.vlm_error = f'{type(e).__name__}: {str(e)[:300]}'
            try:
                import torch
                if torch.cuda.is_available():
                    torch.cuda.empty_cache()
            except Exception:
                pass
            print('AVISO: no se pudo cargar el lector VLM ' + model_name + '. Se continúa con TrOCR como único lector '
                  'manuscrito, así que la confianza ALTA no será alcanzable en esta pasada. Motivo: ' + self.vlm_error)

    def _cargar_paddle(self, paddle_gpu, modo_rapido, cpu_threads, orientacion_auto):
        from paddleocr import PaddleOCR
        paddle_kwargs = dict(use_doc_orientation_classify=bool(orientacion_auto), use_doc_unwarping=False,
                             use_textline_orientation=False, enable_mkldnn=False,
                             cpu_threads=int(cpu_threads or os.cpu_count() or 2))
        if modo_rapido:
            # Al nombrar modelos, PaddleOCR ignora lang/ocr_version: se fijan los dos.
            # lang='es' equivale a PP-OCRv5_server_det + latin_PP-OCRv5_mobile_rec.
            paddle_kwargs.update(text_detection_model_name='PP-OCRv5_mobile_det',
                                 text_recognition_model_name='latin_PP-OCRv5_mobile_rec')
        else:
            paddle_kwargs.update(lang='es', ocr_version='PP-OCRv5')
        self.paddle_device = 'cpu'
        if paddle_gpu:
            try:
                import paddle
                if not paddle.device.is_compiled_with_cuda():
                    raise RuntimeError('paddlepaddle instalado sin CUDA')
                self.paddle = PaddleOCR(device='gpu', **paddle_kwargs)
                self.paddle_device = 'gpu'
            except Exception as e:
                print('PaddleOCR en GPU no disponible; se usará CPU. Motivo:', str(e)[:200])
        if self.paddle is None:
            # CPU para Paddle evita conflictos entre sus bibliotecas CUDA y PyTorch.
            self.paddle = PaddleOCR(device='cpu', **paddle_kwargs)
        self.detector_nombre = 'PaddleOCR PP-OCRv5' + (' mobile-det' if modo_rapido else '') + ' (' + self.paddle_device + ')'

    def _cargar_vlm(self, model_name, gpu):
        """Modelo de visión-lenguaje como lector independiente del recuadro de fecha."""
        import torch
        from transformers import AutoProcessor, Qwen2_5_VLForConditionalGeneration
        self.vlm_device = 'cuda' if gpu and torch.cuda.is_available() else 'cpu'
        dtype = torch.float16 if self.vlm_device == 'cuda' else torch.float32
        self.vlm_processor = AutoProcessor.from_pretrained(model_name, min_pixels=64*28*28, max_pixels=640*28*28)
        self.vlm = Qwen2_5_VLForConditionalGeneration.from_pretrained(model_name, torch_dtype=dtype).to(self.vlm_device).eval()
        self.vlm_prompt = ('Esta imagen es el recuadro "FECHA DE NOTIFICACIÓN" de un acta. Transcribe únicamente la fecha '
                           'manuscrita con el formato DD/MM/AAAA. Si falta el día, el mes o el año, o no es legible, '
                           'responde exactamente ILEGIBLE. No inventes cifras.')

    def read_vlm(self, im):
        import torch
        messages = [{'role': 'user', 'content': [{'type': 'image'}, {'type': 'text', 'text': self.vlm_prompt}]}]
        text = self.vlm_processor.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
        inputs = self.vlm_processor(text=[text], images=[im.convert('RGB')], return_tensors='pt').to(self.vlm_device)
        with torch.inference_mode():
            out = self.vlm.generate(**inputs, max_new_tokens=16, do_sample=False)
        return self.vlm_processor.batch_decode(out[:, inputs['input_ids'].shape[1]:], skip_special_tokens=True)[0].strip()

    def read(self, im, digits=False):
        """Detecciones [{'box','text','score'}] en coordenadas de la imagen recibida."""
        import numpy as np
        if self.paddle is not None:
            # Paddle recibe matrices BGR; PIL entrega RGB.
            output = []
            for result in self.paddle.predict(np.array(im.convert('RGB'))[:, :, ::-1].copy()):
                for box, text, score in zip(result['rec_polys'], result['rec_texts'], result['rec_scores']):
                    output.append({'box': [[float(x), float(y)] for x, y in box],
                                   'text': str(text), 'score': float(score)})
            return output
        kwargs = {'detail': 1, 'paragraph': False, 'batch_size': 1}
        if digits:
            kwargs['allowlist'] = '0123456789/.- '
        rows = self.reader.readtext(np.array(im.convert('RGB')), **kwargs)
        return [{'box': [[float(x), float(y)] for x, y in box], 'text': str(text), 'score': float(score)} for box, text, score in rows]

    def read_handwritten(self, im):
        return self.read_handwritten_lote([im])[0]

    def read_handwritten_lote(self, imagenes):
        """Una sola pasada de TrOCR para varias vistas del mismo recorte."""
        import torch
        pixels = self.processor(images=[im.convert('RGB') for im in imagenes], return_tensors='pt').pixel_values.to(self.trocr_device)
        with torch.inference_mode():
            ids = self.trocr.generate(pixels, max_new_tokens=32, num_beams=4, do_sample=False)
        return [t.strip() for t in self.processor.batch_decode(ids, skip_special_tokens=True)]

    def leer(self, path):
        """Lee una fotografía (ruta o archivo en memoria) y propone los campos.

        La detección de rótulos se hace sobre una copia reducida (rápida); el recuadro de
        la fecha se recorta de la imagen a resolución completa para los lectores manuscritos.
        """
        with Image.open(path) as src:
            if getattr(src, 'n_frames', 1) > 1:
                raise ValueError('Imagen de varias páginas: sepárala antes de procesar.')
            original = ImageOps.exif_transpose(src).convert('RGB')
        W, H = original.size
        im = original.copy(); im.thumbnail((2200, 2800), Image.Resampling.LANCZOS)
        w, h = im.size
        escala = W / w
        head = self.read(im.crop((0, 0, w, int(h * .32))))
        head_text = '\n'.join(l['text'] for l in agrupar_lineas(head))
        exp, code, res = parse_header(head_text)
        nombre, emision_doc = parse_header_extra(head_text)
        offset = int(h * .50)
        foot = self.read(im.crop((0, offset, w, h)))
        foot_abs = [dict(r, box=[[x, y + offset] for x, y in r['box']]) for r in foot]
        tipo = clasificar_documento(head_text, foot_abs, w, h, exp)
        data = dict(expediente=exp, codigo=code, resolucion=res, nombre=nombre, fecha_notificacion=None,
                    fecha_emision=emision_doc, fecha_texto=None, observaciones='')
        evidence = {'encabezado': head, 'pie': foot_abs, 'roi_fecha': None, 'variantes_fecha': [], 'tipo': tipo,
                    'orientacion': 'horizontal' if W > H else 'vertical', 'tamano': [W, H],
                    'motor': self.nombre, 'aviso': 'Puntuaciones OCR no calibradas; no son probabilidades de acierto.'}
        if tipo == 'ANEXO':
            data['observaciones'] = 'Documento anexo (AVISO PREVIO / CEDULÓN): no aporta fecha de notificación para F.'
            evidence['confianza'] = 'SIN_FECHA'; evidence['lectores_coinciden'] = []
            return data, evidence
        roi_s, metodo = localizar_recuadro_fecha(foot_abs, w, h)
        evidence['metodo_recuadro'] = metodo
        if roi_s is None:
            data['observaciones'] = metodo + (' Media página sin cargo: la fecha está en otra foto.' if tipo == 'REC_SIN_CARGO' else ' Completar manualmente desde la foto.')
            evidence['confianza'] = 'SIN_FECHA'; evidence['lectores_coinciden'] = []
            return data, evidence
        roi = tuple(int(v * escala) for v in roi_s)
        evidence['roi_fecha'] = [round(roi[0] / W, 6), round(roi[1] / H, 6), round(roi[2] / W, 6), round(roi[3] / H, 6)]
        crop = original.crop(roi)
        if crop.height < 120:
            factor = max(2, int(120 / max(1, crop.height)) + 1)
            crop = crop.resize((crop.width * factor, crop.height * factor), Image.Resampling.LANCZOS)
        candidates = []
        views = [('color', crop), ('contraste', ImageOps.autocontrast(ImageOps.grayscale(crop)).convert('RGB'))]
        # Tinta azul aislada: elimina el rótulo impreso y las líneas de la tabla, que confunden a TrOCR.
        tinta, fraccion = vista_tinta(crop)
        evidence['tinta_azul'] = round(fraccion, 5)
        if tinta is not None:
            views.append(('tinta', tinta))
        corregidas = {}
        def registrar(vista, texto, detecciones=None):
            ex = extraer_fechas(texto)
            entrada = {'vista': vista, 'texto': texto, 'fechas': ex['exactas'], 'fechas_corregidas': ex['corregidas']}
            if detecciones is not None:
                entrada['detecciones'] = detecciones
            evidence['variantes_fecha'].append(entrada)
            candidates.append(ex['exactas'])
            for cnd in ex['corregidas']:
                corregidas.setdefault(cnd['fecha'], dict(cnd, vista=vista))
        for label, view in views:
            rows = self.read(view, digits=True)
            text = ' '.join(r['text'] for r in sorted(rows, key=lambda r: min(x for x, y in r['box'])))
            registrar(label, text, rows)
        if self.trocr is not None:
            for (label, _), handwritten in zip(views, self.read_handwritten_lote([v for _, v in views])):
                registrar('trocr_' + label, handwritten)
        if self.vlm is not None:
            for label, view in [('vlm_color', crop)] + ([('vlm_tinta', tinta)] if tinta is not None else []):
                try:
                    respuesta = self.read_vlm(view)
                except Exception as e:
                    respuesta = 'ERROR: ' + str(e)[:120]
                registrar(label, '' if 'ILEGIBLE' in respuesta.upper() else respuesta)
                if 'ILEGIBLE' in respuesta.upper():
                    evidence['variantes_fecha'][-1]['texto'] = respuesta
        # Fecha del acta ('Siendo las __ horas, del día __/__/__'). Solo informa y genera alerta si
        # difiere de la propuesta: no entra en las candidatas ni en la confianza (ver alertas()).
        acta = fecha_acta_negativa(foot_abs)
        evidence['fecha_acta'] = acta
        if acta:
            evidence['variantes_fecha'].append({'vista': 'acta_negativa', 'texto': 'fecha del acta (no cuenta para la confianza)', 'fechas': acta})
        nonempty = {x for seq in candidates for x in seq}
        evidence['candidatas_corregidas'] = [corregidas[k] for k in sorted(corregidas)]
        evidence['propuesta_corregida'] = False
        # Una lectura es una propuesta; todas requieren confirmación visual.
        if len(nonempty) == 1:
            data['fecha_notificacion'] = next(iter(nonempty))
            data['fecha_texto'] = ' | '.join(v['texto'] for v in evidence['variantes_fecha'])
            data['observaciones'] = 'Fecha propuesta por OCR. Cotejar visualmente; la coincidencia entre lectores no garantiza exactitud.'
        elif len(nonempty) > 1:
            data['observaciones'] = 'Lecturas de fecha contradictorias: ' + ', '.join(sorted(nonempty))
        elif len(corregidas) == 1:
            cnd = next(iter(corregidas.values()))
            data['fecha_notificacion'] = cnd['fecha']
            data['fecha_texto'] = cnd['texto']
            data['observaciones'] = (f"Candidata corregida ({cnd['regla']}) a partir de '{cnd['texto']}'. "
                                     'Es ambigua: verifica cada cifra en el recorte antes de confirmar.')
            evidence['propuesta_corregida'] = True
        else:
            data['observaciones'] = 'El OCR no obtuvo una fecha completa. Introducirla desde el recorte; no se adivina el año.'
            if len(corregidas) > 1:
                data['observaciones'] += ' Candidatas corregidas distintas: ' + ', '.join(sorted(corregidas)) + '.'
        nivel, coinciden = nivel_confianza(evidence['variantes_fecha'], data['fecha_notificacion'])
        evidence['confianza'] = nivel; evidence['lectores_coinciden'] = coinciden
        return data, evidence

class Proyecto:
    def __init__(self, fotos, excel, salida, hojas=None):
        hojas = tuple(hojas) if hojas else (HOJA,)
        self.fotos, self.excel, self.salida = Path(fotos).expanduser().resolve(), Path(excel).expanduser().resolve(), Path(salida).expanduser().resolve()
        if not self.fotos.is_dir():
            raise ValueError(f'No existe la carpeta de fotos: {self.fotos}')
        if not self.excel.is_file() or self.excel.suffix.lower() != '.xlsx':
            raise ValueError('Selecciona un archivo .xlsx existente.')
        if self.salida == self.fotos or self.fotos in self.salida.parents:
            raise ValueError('La salida debe estar fuera de la carpeta de fotografías.')
        self.salida.mkdir(parents=True,exist_ok=True)
        self.records_dir = self.salida/'registros'
        self.records_dir.mkdir(exist_ok=True)
        self.fingerprint = sha256(self.excel)
        cfg_path = self.salida/'proyecto.json'
        cfg = {'version':VERSION,'excel_sha256':self.fingerprint,'hojas':list(hojas),
               'fecha_emision':FECHA_EMISION.isoformat()}
        if cfg_path.exists():
            previo = json.loads(cfg_path.read_text(encoding='utf-8'))
            # Las carpetas creadas por versiones anteriores no guardaban fecha_emision.
            if any(k in previo and previo[k] != cfg[k] for k in ('excel_sha256','hojas','fecha_emision')):
                raise ValueError('Esta carpeta de resultados corresponde a otro Excel o configuración. Elige otra carpeta de salida.')
        atomic_json(cfg_path,cfg)
        self.cache_path = self.salida/'inventario_cache.json'
        self.index, self.resumen_hojas, self.epoch = build_index(self.excel,hojas)
        self.records = {}
        for p in self.records_dir.glob('*.json'):
            try:
                r = json.loads(p.read_text(encoding='utf-8'))
                self.records[r['id']] = r
            except Exception:
                raise ValueError(f'Registro de avance dañado: {p}. No se ignora para conservar la trazabilidad.') from None

    def save(self,r):
        atomic_json(self.records_dir/f"{r['id']}.json",r)
        self.records[r['id']] = r

    def listar_fotos(self):
        return sorted(p for p in self.fotos.rglob('*') if p.is_file() and p.suffix.lower() in EXTENSIONES and '__MACOSX' not in p.parts)

    def inventario(self, seleccion=None, rehash=False):
        """Identifica cada foto por el SHA-256 de su contenido.

        En Drive el cálculo completo tarda; por eso se guarda inventario_cache.json
        con (tamaño, fecha de modificación) de cada archivo y se reutiliza la huella
        cuando ambos coinciden. Con rehash=True se vuelve a leer todo el contenido.
        """
        paths = self.listar_fotos() if seleccion is None else sorted(set(Path(p).resolve() for p in seleccion))
        for p in paths:
            if self.fotos not in p.resolve().parents or not p.is_file() or p.suffix.lower() not in EXTENSIONES:
                raise ValueError('Archivo seleccionado fuera de la carpeta de fotos o no admitido: '+str(p))
        cache = {}
        if not rehash and self.cache_path.exists():
            try:
                cache = json.loads(self.cache_path.read_text(encoding='utf-8'))
            except Exception:
                cache = {}
        active = defaultdict(list)
        total_bytes, reutilizadas, nuevo_cache = 0, 0, {}
        for p in _progreso(paths, len(paths), 'Inventario'):
            rel = p.relative_to(self.fotos).as_posix()
            st = p.stat()
            previo = cache.get(rel)
            if previo and previo.get('tam') == st.st_size and previo.get('mtime_ns') == st.st_mtime_ns:
                digest = previo['sha256']; reutilizadas += 1
            else:
                digest = sha256(p)
            nuevo_cache[rel] = {'tam': st.st_size, 'mtime_ns': st.st_mtime_ns, 'sha256': digest}
            active[digest].append(rel)
            total_bytes += st.st_size
        if seleccion is None:
            atomic_json(self.cache_path, nuevo_cache)
        else:
            cache.update(nuevo_cache); atomic_json(self.cache_path, cache)
        guardados = 0
        for digest, aliases in active.items():
            r = self.records.get(digest)
            if r is None:
                r = {'id':digest,'estado':'NUEVO','lectura':None,'aprobacion':None,'archivos':aliases}
            elif r.get('archivos') == aliases:
                continue  # sin cambios: no se reescribe el registro en Drive
            else:
                r['archivos'] = aliases
            self.save(r); guardados += 1
        self.active = set(active)
        print(f'{len(paths)} fotos; {len(active)} contenidos únicos; {len(paths)-len(active)} copias idénticas; {total_bytes/1e9:.2f} GB.')
        print(f'Huellas reutilizadas de la caché: {reutilizadas}; registros escritos: {guardados}.')
        print('Hojas:', self.resumen_hojas)
        self.estado()

    def ids_de(self, paths):
        """Identificadores (SHA-256) de las fotos indicadas, según el inventario actual."""
        if not hasattr(self,'active'):
            raise ValueError('Ejecuta primero el inventario.')
        por_ruta = {name: k for k in self.active for name in self.records[k]['archivos']}
        ids = []
        for p in paths:
            rel = Path(p).resolve().relative_to(self.fotos).as_posix()
            if rel not in por_ruta:
                raise ValueError('Foto no inventariada: '+rel)
            if por_ruta[rel] not in ids:
                ids.append(por_ruta[rel])
        return ids

    def pendientes(self, reintentar_errores=False):
        """Identificadores aún sin lectura (NUEVO, y ERROR si se reintenta)."""
        eligible={'NUEVO'} | ({'ERROR'} if reintentar_errores else set())
        return [k for k in sorted(getattr(self,'active',self.records)) if self.records[k]['estado'] in eligible]

    def path(self,r):
        for name in r['archivos']:
            p = (self.fotos/name).resolve()
            if self.fotos in p.parents and p.is_file():
                return p
        raise FileNotFoundError('La fotografía no está en la carpeta actual.')

    def estado(self):
        counts = Counter(self.records[k]['estado'] for k in getattr(self,'active',self.records))
        print('Estados:',dict(counts))
        return counts

    def procesar(self,motor,limite=100,reintentar_errores=False,reprocesar_pendientes=False,solo=None):
        """Lee las fotos pendientes y guarda cada propuesta al instante.

        solo: colección de identificadores a los que se limita la pasada (p. ej. una
        muestra). Cada foto se lee de Drive una sola vez: la huella se calcula sobre
        los mismos bytes que recibe el OCR. Devuelve un resumen con tiempos.
        """
        if not hasattr(self,'active'):
            raise ValueError('Ejecuta primero el inventario.')
        if limite < 1: raise ValueError('El límite debe ser mayor que cero.')
        eligible={'NUEVO'} | ({'ERROR'} if reintentar_errores else set())
        universo=sorted(self.active) if solo is None else [k for k in solo if k in self.active]
        pending=[self.records[k] for k in universo
                 if self.records[k]['estado'] in eligible or
                 (reprocesar_pendientes and self.records[k]['estado'] in {'POR_REVISAR','PENDIENTE'}
                  and self.records[k].get('motor')!=motor.nombre)][:limite]
        resumen={'procesadas':0,'confirmar':0,'discrepancias':0,'errores':0,'segundos':0.0,'detenido':False}
        if pending:
            print(aviso_lectores(motor))
        inicio=time.monotonic()
        barra=_progreso(pending,len(pending),'OCR')
        for r in barra:
            contar=True
            try:
                p=self.path(r)
                datos=p.read_bytes()
                if hashlib.sha256(datos).hexdigest()!=r['id']: raise ValueError('La foto cambió. Repite el inventario.')
                t0=time.monotonic()
                d,evidence=motor.leer(io.BytesIO(datos))
                evidence['segundos']=round(time.monotonic()-t0,2)
                target,issues=evaluar(d,self.index)
                if r.get('lectura') is not None:
                    r.setdefault('historial_ocr',[]).append({k:r.get(k) for k in
                        ('lectura','evidencia_ocr','motor','leido_en','estado')})
                tipo=evidence.get('tipo','REC')
                r.update(lectura=d,estado='ANEXO' if tipo=='ANEXO' else 'POR_REVISAR',destino=target,problemas=issues,
                         evidencia_ocr=evidence,motor=motor.nombre,leido_en=ahora(),error=None,tipo=tipo,
                         confianza=evidence.get('confianza','SIN_FECHA'),alertas=alertas(d,target,evidence),pareja=None)
                self.save(r)
                resumen['discrepancias' if issues else 'confirmar']+=1
            except KeyboardInterrupt:
                print('Detenido. El avance anterior está guardado.')
                resumen['detenido']=True; contar=False
                break
            except Exception as e:
                r.update(estado='ERROR',error=str(e)[:500]); self.save(r)
                resumen['errores']+=1
                mensaje=f'Error registrado en {r["archivos"][0]}: {str(e)[:120]}'
                (tqdm.write if tqdm is not None else print)(mensaje)
                if 'out of memory' in str(e).lower():
                    print('Memoria insuficiente. Detenido; reinicia el entorno o usa CPU con EasyOCR.')
                    resumen['detenido']=True
                    break
            finally:
                resumen['procesadas']+=contar
                if tqdm is not None and hasattr(barra,'set_postfix'):
                    barra.set_postfix(ok=resumen['confirmar'],rev=resumen['discrepancias'],err=resumen['errores'],refresh=False)
        resumen['segundos']=round(time.monotonic()-inicio,1)
        resumen['emparejadas']=self.emparejar()
        if resumen['procesadas']:
            por_foto=resumen['segundos']/resumen['procesadas']
            restantes=len(self.pendientes(reintentar_errores))
            print(f"{resumen['procesadas']} fotos en {resumen['segundos']/60:.1f} min ({por_foto:.1f} s/foto). "
                  f"Pendientes: {restantes}; estimado {restantes*por_foto/3600:.1f} h a este ritmo.")
        self.estado()
        return resumen

    @staticmethod
    def _orden_natural(nombre):
        return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', nombre)]

    def emparejar(self, distancia=2):
        """Une medias páginas: un cargo sin encabezado toma expediente y código de la REC
        sin cargo inmediatamente anterior en la misma carpeta (orden natural de nombres).
        El emparejamiento es una propuesta que el revisor confirma viendo ambas fotos.
        """
        por_carpeta=defaultdict(list)
        for k in getattr(self,'active',self.records):
            r=self.records[k]
            if r.get('tipo') in ('REC_SIN_CARGO','CARGO_SIN_ENCABEZADO'):
                for nombre in r['archivos']:
                    por_carpeta[posixpath.dirname(nombre)].append((self._orden_natural(posixpath.basename(nombre)),nombre,k))
        nuevas=0
        for carpeta,items in por_carpeta.items():
            items.sort()
            for i,(_,nombre,k) in enumerate(items):
                r=self.records[k]
                if r.get('tipo')!='CARGO_SIN_ENCABEZADO' or r.get('pareja') or r['estado']!='POR_REVISAR':
                    continue
                for j in range(i-1,max(-1,i-1-distancia),-1):
                    _,nombre_prev,kp=items[j]
                    rp=self.records[kp]
                    if rp.get('tipo')=='REC_SIN_CARGO' and not rp.get('pareja'):
                        d=dict(r.get('lectura') or {}); lp=rp.get('lectura') or {}
                        d.update(expediente=lp.get('expediente'),codigo=lp.get('codigo'),resolucion=lp.get('resolucion'))
                        d['observaciones']=(f'Expediente y código tomados de la media página {nombre_prev}. '
                                            'Confirma que ambas fotos son el mismo documento. '+(d.get('observaciones') or ''))
                        target,issues=evaluar(d,self.index)
                        r.update(lectura=d,destino=target,problemas=issues,pareja=kp,alertas=alertas(d,target,r.get('evidencia_ocr')))
                        rp.update(pareja=k,problemas=[f'Media página sin cargo; su fecha se toma de {nombre}.'])
                        self.save(r); self.save(rp); nuevas+=1
                        break
        if nuevas:
            print(f'Medias páginas emparejadas: {nuevas}. Confírmalas en la revisión viendo ambas fotos.')
        return nuevas

    CATEGORIAS = (('AVISO PREVIO', 'AVISO_PREVIO'), ('CEDULON', 'CEDULON'), ('CERRADO', 'DOMICILIO_CERRADO'),
                  ('NEGATIVA', 'NEGATIVA'), ('FIRMA', 'FIRMA'))

    @classmethod
    def categoria_carpeta(cls, rel):
        carpeta = norm(posixpath.dirname(rel)).upper()
        for clave, nombre in cls.CATEGORIAS:
            if norm(clave).upper() in carpeta:
                return nombre
        return 'OTRA'

    def sondeo_formato(self, guardar=True):
        """Recorre todas las fotos sin OCR: orientación, tamaño, formato y categoría de carpeta.

        Sirve para saber si todos los documentos comparten la estructura esperada (REC vertical
        de página completa) antes de leerlos. Solo lee la cabecera de cada archivo.
        """
        filas = []
        paths = self.listar_fotos()
        for p in _progreso(paths, len(paths), 'Sondeo'):
            rel = p.relative_to(self.fotos).as_posix()
            try:
                with Image.open(p) as im:
                    w, h = im.size; fmt = im.format
                    orientacion_exif = im.getexif().get(0x0112, 1)
                if orientacion_exif in (5, 6, 7, 8):
                    w, h = h, w
                fila = {'archivo': rel, 'ancho': w, 'alto': h, 'formato': fmt,
                        'orientacion': 'horizontal' if w > h else 'vertical'}
            except Exception as e:
                fila = {'archivo': rel, 'ancho': None, 'alto': None, 'formato': 'ERROR', 'orientacion': 'ERROR', 'error': str(e)[:120]}
            fila['categoria'] = self.categoria_carpeta(rel)
            filas.append(fila)
        res = {'generado_en': ahora(), 'fotos': len(filas),
               'por_categoria_y_orientacion': {f'{k[0]} {k[1]}': v for k, v in sorted(Counter((f['categoria'], f['orientacion']) for f in filas).items())},
               'por_tamano': {f'{f[0]}x{f[1]}': v for f, v in Counter((f['ancho'], f['alto']) for f in filas).most_common()},
               'por_formato': dict(Counter(f['formato'] for f in filas)),
               'errores': [f for f in filas if f['formato'] == 'ERROR'], 'detalle': filas}
        print(f"Sondeo de formato: {len(filas)} fotos; vertical {sum(1 for f in filas if f['orientacion']=='vertical')}, "
              f"horizontal {sum(1 for f in filas if f['orientacion']=='horizontal')}, errores {len(res['errores'])}.")
        for k, v in res['por_categoria_y_orientacion'].items(): print(f'  {k:<32}{v:>6}')
        print('  Tamaños más frecuentes:', ', '.join(f'{k}: {v}' for k, v in list(res['por_tamano'].items())[:8]))
        if guardar:
            ruta = self.salida / f"sondeo_formato_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
            atomic_json(ruta, res); res['archivo'] = str(ruta); print('Guardado en', ruta)
        return res

    CAMPOS_CLAVE = ('expediente', 'nombre', 'codigo', 'fecha_emision', 'fecha_notificacion')

    def sondeo_estructura(self, ids=None, guardar=True):
        """Sobre fotos ya leídas: tipo de documento, campos clave encontrados y cómo se ubicó la fecha.

        'Estructura completa' = REC con expediente, nombre, código y emisión leídos y recuadro de fecha localizado.
        """
        ids = [k for k in (ids if ids is not None else sorted(getattr(self, 'active', self.records))) if self.records[k].get('lectura') is not None]
        filas = []
        for k in ids:
            r = self.records[k]; d = r.get('lectura') or {}; ev = r.get('evidencia_ocr') or {}
            encontrados = {c: bool(d.get(c)) for c in self.CAMPOS_CLAVE}
            filas.append({'archivo': r['archivos'][0], 'categoria': self.categoria_carpeta(r['archivos'][0]), 'tipo': r.get('tipo') or '-',
                          'orientacion': ev.get('orientacion'), 'recuadro': ev.get('metodo_recuadro') or ('no' if ev.get('roi_fecha') is None else 'si'),
                          'campos': encontrados, 'confianza': r.get('confianza'),
                          'completa': r.get('tipo') == 'REC' and all(encontrados[c] for c in ('expediente', 'nombre', 'codigo', 'fecha_emision')) and ev.get('roi_fecha') is not None})
        n = len(filas)
        res = {'generado_en': ahora(), 'fotos_leidas': n,
               'por_tipo': dict(Counter(f['tipo'] for f in filas)),
               'por_categoria_y_tipo': {f'{k[0]} {k[1]}': v for k, v in sorted(Counter((f['categoria'], f['tipo']) for f in filas).items())},
               'campos_encontrados': {c: sum(1 for f in filas if f['campos'][c]) for c in self.CAMPOS_CLAVE},
               'recuadro_fecha': dict(Counter(f['recuadro'] for f in filas)),
               'estructura_completa': sum(1 for f in filas if f['completa']),
               'incompletas': [f for f in filas if not f['completa']], 'detalle': filas}
        if n:
            print(f'Sondeo de estructura sobre {n} fotos leídas. Estructura completa: {res["estructura_completa"]} ({100*res["estructura_completa"]/n:.0f}%).')
            print('  Tipos:', res['por_tipo'])
            print('  Campos encontrados:', ', '.join(f'{c} {v} ({100*v/n:.0f}%)' for c, v in res['campos_encontrados'].items()))
            print('  Recuadro de fecha:', res['recuadro_fecha'])
            for k, v in res['por_categoria_y_tipo'].items(): print(f'  {k:<40}{v:>6}')
        else:
            print('No hay fotos leídas todavía; ejecuta el piloto o el procesamiento.')
        if guardar:
            ruta = self.salida / f"sondeo_estructura_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
            atomic_json(ruta, res); res['archivo'] = str(ruta); print('Guardado en', ruta)
        return res

    def conciliar(self, guardar=True):
        """Filas del Excel en el alcance (hoja y E = lote) frente a las fotos, por expediente.

        Antes del OCR solo compara cantidades; con las fotos leídas indica qué filas tienen foto,
        cuáles no, qué fotos apuntan a expedientes fuera del lote y cuáles no traen expediente.
        """
        alcance = {k: items[0] for k, items in self.index.items() if len(items) == 1 and items[0]['hoja'] == HOJA and fecha(items[0]['emision']) == FECHA_EMISION}
        activos = sorted(getattr(self, 'active', self.records))
        por_exp = defaultdict(list)
        sin_expediente, no_leidas, anexos = [], [], []
        for k in activos:
            r = self.records[k]
            if r['estado'] == 'NUEVO':
                no_leidas.append(k); continue
            if r['estado'] == 'ANEXO':
                anexos.append(k); continue
            parsed = expediente((r.get('lectura') or {}).get('expediente'))
            if parsed and parsed[1]:
                por_exp[parsed].append(k)
            else:
                sin_expediente.append(k)
        con_foto = sorted(k for k in alcance if k in por_exp)
        sin_foto = [alcance[k] for k in sorted(alcance) if k not in por_exp]
        fuera = {f'{k[0]}-{k[1]}': [self.records[i]['archivos'][0] for i in ids] for k, ids in sorted(por_exp.items()) if k not in alcance}
        multiples = {f'{k[0]}-{k[1]}': [self.records[i]['archivos'][0] for i in por_exp[k]] for k in con_foto if len(por_exp[k]) > 1}
        res = {'generado_en': ahora(), 'alcance': {'hoja': HOJA, 'fecha_emision': FECHA_EMISION.isoformat()},
               'filas_en_alcance': len(alcance), 'filas_f_vacia': sum(1 for v in alcance.values() if v['fecha_actual'] in (None, '')),
               'fotos_archivos': sum(len(self.records[k]['archivos']) for k in activos), 'fotos_unicas': len(activos),
               'fotos_no_leidas': len(no_leidas), 'fotos_anexo': len(anexos), 'fotos_sin_expediente': len(sin_expediente),
               'filas_con_foto': len(con_foto), 'filas_sin_foto': len(sin_foto),
               'expedientes_fuera_del_lote': len(fuera), 'filas_con_varias_fotos': len(multiples),
               'detalle': {'filas_sin_foto': [{'expediente': v['expediente'], 'celda': v['celda'], 'f_actual': v['fecha_actual']} for v in sin_foto],
                           'fuera_del_lote': fuera, 'varias_fotos': multiples,
                           'sin_expediente': [self.records[k]['archivos'][0] for k in sin_expediente]}}
        print(f"Conciliación: {len(alcance)} filas con E = {FECHA_EMISION:%d/%m/%Y} en la hoja {HOJA} ({res['filas_f_vacia']} con F vacía) "
              f"frente a {res['fotos_archivos']} fotos ({len(activos)} únicas).")
        if no_leidas:
            print(f'  Fotos aún no leídas: {len(no_leidas)}. La comparación por expediente es parcial hasta terminar el procesamiento.')
        print(f"  Filas con foto: {len(con_foto)} | filas sin foto: {len(sin_foto)} | filas con varias fotos: {len(multiples)}")
        print(f"  Fotos con expediente fuera del lote: {len(fuera)} expedientes | fotos sin expediente legible: {len(sin_expediente)} | anexos: {len(anexos)}")
        if guardar:
            ruta = self.salida / f"conciliacion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
            atomic_json(ruta, res); res['archivo'] = str(ruta); print('Guardado en', ruta)
        return res

    def resumen_alcance(self):
        """Filas de la hoja dentro del alcance (E = fecha de emisión) con F vacía y con F llena."""
        vacias=llenas=0
        for items in self.index.values():
            for r in items:
                if r['hoja']==HOJA and fecha(r['emision'])==FECHA_EMISION:
                    if r['fecha_actual'] in (None,''): vacias+=1
                    else: llenas+=1
        print(f'Hoja {HOJA}, E = {FECHA_EMISION:%d/%m/%Y}: {vacias+llenas} filas; F vacía: {vacias}; F ya registrada: {llenas}.')
        return {'filas':vacias+llenas,'f_vacia':vacias,'f_registrada':llenas}

    def medir_acierto(self, ids=None, guardar=True, minimo_alta=30):
        """Compara la fecha propuesta por el OCR con la columna F ya registrada en el Excel.

        Solo usa fotos cuyo expediente coincide con una fila que ya tiene F: son el patrón
        de referencia para estimar el acierto real antes del lote completo. Imprime la tabla
        confianza x resultado y guarda piloto_<marca>.json en la carpeta de resultados.
        Con menos de minimo_alta casos ALTA con referencia, el acierto de ALTA no es concluyente.
        """
        ids=list(ids) if ids is not None else sorted(getattr(self,'active',self.records))
        res={'generado_en':ahora(),'version':VERSION,'fotos_evaluadas':len(ids),'con_referencia':0,'aciertos':0,
             'errores':0,'sin_fecha_ocr':0,'detalle':[]}
        for k in ids:
            r=self.records[k]; d=r.get('lectura') or {}
            parsed=expediente(d.get('expediente'))
            if not parsed or parsed[1] is None: continue
            matches=self.index.get(parsed,[])
            if len(matches)!=1 or matches[0]['fecha_actual'] in (None,''): continue
            ref=fecha(matches[0]['fecha_actual']); ocr=fecha(d.get('fecha_notificacion'))
            res['con_referencia']+=1
            if ocr is None: res['sin_fecha_ocr']+=1; estado='SIN_FECHA'
            elif ocr==ref: res['aciertos']+=1; estado='ACIERTO'
            else: res['errores']+=1; estado='ERROR'
            res['detalle'].append({'archivo':r['archivos'][0],'excel':ref.isoformat() if ref else None,
                                   'ocr':ocr.isoformat() if ocr else None,'confianza':r.get('confianza'),'resultado':estado})
        n=res['con_referencia']
        niveles=['ALTA','MEDIA','CONTRADICTORIA','SIN_FECHA']; resultados=['ACIERTO','ERROR','SIN_FECHA']
        tabla={}
        for nivel in niveles+['TOTAL']:
            filas=[x for x in res['detalle'] if nivel=='TOTAL' or (x['confianza'] or 'SIN_FECHA')==nivel]
            total=len(filas)
            tabla[nivel]={'total':total}
            for est in resultados:
                k=sum(1 for x in filas if x['resultado']==est)
                tabla[nivel][est]={'n':k,'pct':round(100*k/total,1) if total else None}
        res['tabla']=tabla
        res['alta_con_referencia']=tabla['ALTA']['total']
        res['concluyente']=tabla['ALTA']['total']>=minimo_alta
        res['tipos']=dict(Counter(self.records[k].get('tipo') or '-' for k in ids))
        res['confianza_muestra']=dict(Counter(self.records[k].get('confianza') or (self.records[k]['estado'] if self.records[k]['estado'] in ('ERROR','ANEXO') else 'SIN_LECTURA') for k in ids))
        if n:
            print(f"Fotos con fecha de referencia en el Excel: {n} de {len(ids)}.")
            print(f"{'confianza':<15}{'total':>6}  "+'  '.join(f'{e:>16}' for e in resultados))
            for nivel in niveles+['TOTAL']:
                fila=tabla[nivel]
                celdas='  '.join(f"{fila[e]['n']:>5} ({fila[e]['pct'] if fila[e]['pct'] is not None else 0:>5.1f}%)" for e in resultados)
                print(f"{nivel:<15}{fila['total']:>6}  {celdas}")
            for x in res['detalle']:
                if x['resultado']=='ERROR': print('  ERROR', x['archivo'], 'Excel', x['excel'], 'OCR', x['ocr'], 'confianza', x['confianza'])
            if not res['concluyente']:
                print(f"AVISO: solo {tabla['ALTA']['total']} casos ALTA con referencia (mínimo {minimo_alta}); el acierto de ALTA no es concluyente. Amplía la muestra.")
        else:
            print('Ninguna foto de la muestra corresponde a una fila con F ya registrada; amplía la muestra.')
        if guardar:
            ruta=self.salida/f"piloto_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
            atomic_json(ruta,res); res['archivo']=str(ruta); print('Piloto guardado en',ruta)
        return res

    def aprobar(self,rid,data,revisor,verifico_rec_uno=False):
        """Registra la confirmación visual de un revisor.

        verifico_rec_uno: obligatorio cuando no se leyó el número de resolución; el revisor
        declara haber comprobado en la foto que se trata de la REC UNO. Si se leyó otra
        resolución, evaluar() bloquea la aprobación.
        """
        if not revisor.strip():
            raise ValueError('Escribe el nombre o iniciales de quien revisa.')
        r = self.records[rid]
        if sha256(self.path(r)) != rid:
            raise ValueError('La imagen cambió. Ejecuta el inventario otra vez.')
        data = normalizar_lectura(data)
        target,issues = evaluar(data,self.index)
        if issues:
            raise ValueError('\n'.join(issues))
        if not target:
            raise ValueError('Falta el destino.')
        if data.get('resolucion') in (None,'') and not verifico_rec_uno:
            raise ValueError('No se leyó el número de resolución. Marca «Verifiqué visualmente que es la REC UNO» para aprobar, o escribe la resolución leída en la foto.')
        previous = r.get('aprobacion')
        if previous:
            mismos = all((previous.get('datos') or {}).get(f)==data.get(f) for f in ('expediente','codigo','resolucion','fecha_notificacion'))
            if mismos and any(a['revisor']==revisor.strip() for a in self.aprobaciones_vigentes(r)):
                raise ValueError('Este revisor ya confirmó estos mismos datos. La segunda revisión debe hacerla otra persona.')
            r.setdefault('historial_aprobaciones',[]).append(previous)
        entrada = {'datos':data,'destino':target,'revisor':revisor.strip(),'fecha':ahora(),
                   'confianza_ocr':r.get('confianza'),'alertas':alertas(data,target,r.get('evidencia_ocr')),
                   'verifico_rec_uno':bool(verifico_rec_uno),
                   'lectura_ocr':(r.get('lectura') or {}).get('fecha_notificacion'),
                   'evidencia':self.guardar_evidencia(r)}
        r.setdefault('aprobaciones',[]).append(entrada)
        r['aprobacion'] = entrada
        r['estado'] = 'APROBADO'
        self.save(r)

    def aprobaciones_vigentes(self, r):
        """Aprobaciones cuyos datos coinciden con la última (una corrección posterior las invalida)."""
        ultima = r.get('aprobacion')
        if not ultima:
            return []
        clave = lambda a: tuple((a.get('datos') or {}).get(k) for k in ('expediente','codigo','resolucion','fecha_notificacion'))
        return [a for a in r.get('aprobaciones') or [ultima] if clave(a) == clave(ultima)]

    def guardar_evidencia(self, r):
        """Guarda en Drive el cargo y el recorte de fecha que vio el revisor, nombrados por la huella de la foto."""
        carpeta = self.salida/'evidencia'
        carpeta.mkdir(exist_ok=True)
        rutas = {}
        with Image.open(self.path(r)) as src:
            im = ImageOps.exif_transpose(src).convert('RGB')
        w,h = im.size
        cargo = im.crop((0,int(h*.58),w,h)); cargo.thumbnail((1600,1600))
        destino = carpeta/f"{r['id']}_cargo.jpg"; cargo.save(destino,'JPEG',quality=90); rutas['cargo']=destino.name
        roi = (r.get('evidencia_ocr') or {}).get('roi_fecha')
        if roi:
            rec = im.crop((int(roi[0]*w),int(roi[1]*h),int(roi[2]*w),int(roi[3]*h)))
            rec = rec.resize((rec.width*2,rec.height*2),Image.Resampling.LANCZOS)
            destino = carpeta/f"{r['id']}_fecha.jpg"; rec.save(destino,'JPEG',quality=95); rutas['fecha']=destino.name
        return rutas

    def revision_manual(self, limite=100):
        """Permite trabajar aun sin modelos descargados ni OCR disponible."""
        pending=[self.records[k] for k in sorted(self.active) if self.records[k]['estado']=='NUEVO'][:limite]
        for r in pending:
            r.update(estado='PENDIENTE',lectura=None)
            self.save(r)
        print(f'{len(pending)} fotografías preparadas para transcripción visual.')

    def lista_revision(self, modo='pendientes'):
        """Identificadores a revisar: 'pendientes', 'segunda' (aprobados por un solo revisor) o 'todos'."""
        ids = sorted(getattr(self,'active',self.records), key=lambda k:(self.records[k]['archivos'][0]))
        estados = {'pendientes':{'POR_REVISAR','PENDIENTE','ERROR'},'segunda':{'APROBADO'},
                   'todos':{'POR_REVISAR','PENDIENTE','ERROR','APROBADO','ANEXO'}}[modo]
        ids = [k for k in ids if self.records[k]['estado'] in estados]
        if modo == 'segunda':
            ids = [k for k in ids if len({a['revisor'] for a in self.aprobaciones_vigentes(self.records[k])}) < 2]
        orden = {'ALTA':0,'MEDIA':1,'CONTRADICTORIA':2,'SIN_FECHA':3}
        def prioridad(k):
            r=self.records[k]; d=r.get('lectura') or {}
            t,_=evaluar(d,self.index)
            en_alcance = 0 if t and fecha(t['emision'])==FECHA_EMISION and t['fecha_actual'] in (None,'') else 1
            return (en_alcance, orden.get(r.get('confianza'),3))
        ids.sort(key=prioridad)
        return ids

    def revisar(self, modo='pendientes'):
        import ipywidgets as W
        from IPython.display import display, clear_output
        ids = self.lista_revision(modo)
        if not ids:
            print('No hay registros para este modo. Ejecuta el procesamiento o cambia el modo de revisión.')
            return
        def etiqueta(k):
            r=self.records[k]
            revisores=sorted({a['revisor'] for a in self.aprobaciones_vigentes(r)})
            return f'{r["estado"]} | {r.get("confianza") or "-"} | {r.get("tipo") or "-"} | {r["archivos"][0]}'+(f' | revisado por {", ".join(revisores)}' if revisores else '')
        choice = W.Dropdown(options=[(etiqueta(k),k) for k in ids],layout=W.Layout(width='98%'))
        output, msg = W.Output(), W.Output()
        exp = W.Text(description='Expediente:',layout=W.Layout(width='98%'),style={'description_width':'140px'})
        nom = W.Text(description='Nombre leído:',layout=W.Layout(width='98%'),style={'description_width':'140px'})
        code = W.Text(description='Código:',style={'description_width':'140px'})
        res = W.Text(description='Resolución:',style={'description_width':'140px'})
        dt = W.Text(description='Notificación:',placeholder='DD/MM/AAAA',style={'description_width':'140px'})
        em = W.Text(description='Emisión REC:',placeholder='DD/MM/AAAA',style={'description_width':'140px'})
        who = W.Text(description='Revisor:',style={'description_width':'140px'})
        check = W.Checkbox(value=False,description='He cotejado expediente, código, resolución y fecha con la foto.',indent=False,layout=W.Layout(width='98%'))
        rec_uno = W.Checkbox(value=False,description='Verifiqué visualmente que es la REC UNO (obligatorio si la resolución no se leyó).',indent=False,layout=W.Layout(width='98%'))
        approve = W.Button(description='Confirmar y guardar',button_style='success')
        skip = W.Button(description='Dejar pendiente')
        next_btn = W.Button(description='Siguiente')
        rotate = W.Button(description='Girar vista 90°')
        full = W.Checkbox(value=False,description='Mostrar página completa')
        angle = [0]
        def show_image():
            with output:
                clear_output(wait=True)
                r = self.records[choice.value]
                with Image.open(self.path(r)) as src:
                    im = ImageOps.exif_transpose(src).convert('RGB').rotate(angle[0],expand=True)
                w,h = im.size
                for title,view in [('Encabezado',im.crop((0,0,w,int(h*.30)))),('Cargo de notificación',im.crop((0,int(h*.62),w,h)))]:
                    print(title)
                    view.thumbnail((1100,1000))
                    display(view)
                roi=(r.get('evidencia_ocr') or {}).get('roi_fecha')
                if roi and angle[0]==0:
                    crop=im.crop((int(roi[0]*w),int(roi[1]*h),int(roi[2]*w),int(roi[3]*h)))
                    crop=crop.resize((crop.width*2,crop.height*2),Image.Resampling.LANCZOS)
                    crop.thumbnail((1100,350))
                    print('Recorte que leyó el OCR: comprueba que contenga la fecha correcta.')
                    display(crop)
                    tinta,_=vista_tinta(crop)
                    if tinta is not None:
                        print('Solo la tinta azul (manuscrito aislado):'); display(tinta)
                if full.value:
                    im.thumbnail((1100,1700)); display(im)
                if r.get('pareja') and r['pareja'] in self.records:
                    rp=self.records[r['pareja']]
                    print('Media página emparejada:',rp['archivos'][0])
                    with Image.open(self.path(rp)) as src2:
                        im2=ImageOps.exif_transpose(src2).convert('RGB')
                    im2.thumbnail((1100,900)); display(im2)
                data = r.get('aprobacion',{}).get('datos') if r.get('aprobacion') else r.get('lectura')
                if data:
                    target,issues = evaluar(data,self.index)
                    print('Destino:',target)
                    if target: print('Nombre en la fila del Excel:',target.get('nombre') or '-','| emisión impresa leída:',data.get('fecha_emision') or '-')
                    print('Confianza OCR:',r.get('confianza') or '-','| lectores que coinciden:',', '.join((r.get('evidencia_ocr') or {}).get('lectores_coinciden') or []) or 'ninguno')
                    for v in (r.get('evidencia_ocr') or {}).get('variantes_fecha') or []:
                        print(f"  {v['vista']:<12} -> {', '.join(v.get('fechas') or []) or 'sin fecha'} | texto: {str(v.get('texto'))[:60]}")
                    for cnd in (r.get('evidencia_ocr') or {}).get('candidatas_corregidas') or []:
                        print(f"  candidata corregida ({cnd['regla']}, {cnd.get('vista')}): {cnd['fecha']} desde '{cnd['texto']}'. No cuenta para la confianza.")
                    print('Observaciones de lectura:',data.get('observaciones') or 'Ninguna')
                    print('Comprobaciones:',issues or 'Coinciden; falta tu comprobación visual.')
                    for aviso in alertas(data,target,r.get('evidencia_ocr')): print('ALERTA:',aviso)
                    vigentes=self.aprobaciones_vigentes(r)
                    if vigentes:
                        print('Aprobaciones vigentes:',', '.join(f"{a['revisor']} ({a['fecha']})" for a in vigentes))
                elif r.get('error'):
                    print('Error:',r['error'])
        def load(change=None):
            r = self.records[choice.value]
            d = (r['aprobacion']['datos'] if r.get('aprobacion') else r.get('lectura')) or {}
            for widget,field in [(exp,'expediente'),(nom,'nombre'),(code,'codigo'),(res,'resolucion'),(dt,'fecha_notificacion'),(em,'fecha_emision')]:
                widget.value = d.get(field) or ''
            angle[0]=0; check.value=False; rec_uno.value=False
            with msg: clear_output()
            show_image()
        def advance(_=None):
            idx = ids.index(choice.value)
            if idx+1 < len(ids): choice.value = ids[idx+1]
            else:
                with msg: print('Fin de la lista. Puedes seleccionar otro registro arriba.')
        def confirm(_):
            with msg:
                clear_output()
                try:
                    if not check.value: raise ValueError('Marca la casilla después de cotejar la fotografía.')
                    r = self.records[choice.value]
                    d = dict(r.get('lectura') or {})
                    d.update(expediente=exp.value,nombre=nom.value,codigo=code.value,resolucion=res.value,
                             fecha_notificacion=dt.value,fecha_emision=em.value)
                    self.aprobar(choice.value,d,who.value,verifico_rec_uno=rec_uno.value)
                    n=len({a['revisor'] for a in self.aprobaciones_vigentes(self.records[choice.value])})
                    print(f'Confirmado y guardado ({n} revisor(es) coinciden). Pulsa Siguiente para continuar.')
                    check.value=False
                except Exception as e: print(str(e))
        def postpone(_):
            r=self.records[choice.value]
            if r.get('aprobacion'): r.setdefault('historial_aprobaciones',[]).append(r['aprobacion'])
            r.update(estado='PENDIENTE',aprobacion=None)
            self.save(r)
            with msg: clear_output(); print('Pendiente: esta foto no se exportará al Excel.')
        def turn(_):
            angle[0]=(angle[0]+90)%360
            show_image()
        choice.observe(load,names='value'); approve.on_click(confirm); skip.on_click(postpone)
        next_btn.on_click(advance); rotate.on_click(turn); full.observe(lambda _:show_image(),names='value')
        display(W.VBox([choice,full,rotate,output,exp,nom,code,res,dt,em,who,check,rec_uno,W.HBox([approve,skip,next_btn]),msg]))
        load()

    def exportar(self, doble_revision=False):
        """Escribe las fechas aprobadas en una copia del Excel y deja constancia.

        doble_revision: False exporta toda aprobación válida; 'no_alta' exige dos revisores
        distintos cuando la confianza OCR no fue ALTA; True exige dos revisores siempre.
        Genera además revision_*.csv y acta_exportacion_*.json con las huellas SHA-256.
        """
        if sha256(self.excel) != self.fingerprint:
            raise ValueError('El Excel original cambió. Usa una carpeta de resultados nueva y vuelve a verificar las coincidencias.')
        if doble_revision not in (False,True,'no_alta'):
            raise ValueError("doble_revision debe ser False, True o 'no_alta'.")
        proposals = defaultdict(list)
        report = []
        for rid in sorted(getattr(self,'active',self.records)):
            r = self.records[rid]
            a = r.get('aprobacion') if r['estado']=='APROBADO' else None
            d = (a or {}).get('datos') or r.get('lectura') or {}
            target,issues = evaluar(d,self.index)
            vigentes = self.aprobaciones_vigentes(r) if a else []
            revisores = sorted({x['revisor'] for x in vigentes})
            entry = {'id':rid,'archivos':' | '.join(r['archivos']),'estado':r['estado'],
                     'expediente':d.get('expediente'),'codigo':d.get('codigo'),'fecha':d.get('fecha_notificacion'),
                     'nombre_doc':d.get('nombre') or '','nombre_excel':(target or {}).get('nombre') or '','emision_doc':d.get('fecha_emision') or '',
                     'hoja':target['hoja'] if target else '', 'celda':target['celda'] if target else '',
                     'revisor':' | '.join(revisores), 'revisiones':len(revisores),
                     'aprobado_en':a['fecha'] if a else '', 'confianza_ocr':r.get('confianza') or '','tipo':r.get('tipo') or '',
                     'verifico_rec_uno':('si' if a.get('verifico_rec_uno') else 'no') if a else '',
                     'lectura_ocr':(r.get('lectura') or {}).get('fecha_notificacion') or '',
                     'alertas':' | '.join(alertas(d,target,r.get('evidencia_ocr'))),
                     'evidencia':' | '.join(sorted((a.get('evidencia') or {}).values())) if a else '',
                     'resultado':'NO EXPORTADO',
                     'detalle':'; '.join(issues) or r.get('error') or ''}
            report.append(entry)
            if a and not issues and target:
                if sha256(self.path(r)) != rid:
                    entry['detalle']='La fotografía cambió desde la aprobación.'
                    continue
                exige_dos = doble_revision is True or (doble_revision=='no_alta' and r.get('confianza')!='ALTA')
                if exige_dos and len(revisores) < 2:
                    entry.update(resultado='FALTA SEGUNDA REVISION',detalle='Requiere confirmación de un segundo revisor distinto.')
                    continue
                proposals[(target['hoja'],target['celda'])].append((fecha(d['fecha_notificacion']),entry))
        updates = {}
        for key, items in proposals.items():
            ds = {d for d,_ in items}
            if len(ds)>1:
                for _,entry in items:
                    entry.update(resultado='CONFLICTO',detalle='Dos fotografías aprobadas proponen fechas distintas para la misma celda.')
                continue
            updates[key]=next(iter(ds))
            for _,entry in items: entry['resultado']='INCLUIDO'
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:6]
        report_path = self.salida/f'revision_{stamp}.csv'
        columns = ['id','archivos','estado','tipo','expediente','codigo','fecha','nombre_doc','nombre_excel','emision_doc','hoja','celda','revisor','revisiones','aprobado_en',
                   'verifico_rec_uno','confianza_ocr','lectura_ocr','alertas','evidencia','resultado','detalle']
        def safe(v):
            s = str(v if v is not None else '')
            return "'"+s if s.startswith(('=','+','-','@','\t','\r')) else s
        with report_path.open('w',newline='',encoding='utf-8-sig') as f:
            writer=csv.DictWriter(f,fieldnames=columns,delimiter=';')
            writer.writeheader()
            writer.writerows({k:safe(v) for k,v in r.items()} for r in report)
        if not updates:
            print('Todavía no hay fechas aprobadas sin conflictos. Informe:',report_path)
            return None,report_path
        output = self.salida/f'expedientes_actualizados_{stamp}.xlsx'
        patch_xlsx(self.excel,output,updates,self.epoch)
        acta = {'version':VERSION,'generado_en':ahora(),'excel_original':{'ruta':str(self.excel),'sha256':self.fingerprint},
                'excel_generado':{'ruta':str(output),'sha256':sha256(output)},'informe_csv':report_path.name,
                'alcance':{'hoja':HOJA,'fecha_emision':FECHA_EMISION.isoformat()},'doble_revision':doble_revision,
                'celdas_escritas':len(updates),
                'detalle':[{k:e[k] for k in ('hoja','celda','fecha','expediente','codigo','id','archivos','revisor','revisiones',
                                             'aprobado_en','verifico_rec_uno','confianza_ocr','lectura_ocr','evidencia')}
                           for e in report if e['resultado']=='INCLUIDO']}
        atomic_json(self.salida/f'acta_exportacion_{stamp}.json',acta)
        print(f'Excel generado: {len(updates)} celdas confirmadas. Las demás conservan su valor anterior.')
        print(output); print(report_path); print(self.salida/f'acta_exportacion_{stamp}.json')
        return output, report_path

    def paquete_revision(self, limite=20):
        """Exporta evidencia y lecturas para adjuntarlas a ChatGPT; no llama ninguna API."""
        ids=[k for k in sorted(getattr(self,'active',self.records)) if self.records[k]['estado'] in {'POR_REVISAR','PENDIENTE','ERROR'}][:limite]
        if not ids: print('No hay casos pendientes.'); return None
        stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
        dest=self.salida/f'revision_chatgpt_{stamp}.zip'
        manifest=[]
        with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
            for i,k in enumerate(ids,1):
                r=self.records[k]
                with Image.open(self.path(r)) as src:
                    im=ImageOps.exif_transpose(src).convert('RGB')
                w,h=im.size
                for name,view in [('pagina',im),('encabezado',im.crop((0,0,w,int(h*.30)))),('cargo',im.crop((0,int(h*.62),w,h)))]:
                    view.thumbnail((2000,2400))
                    b=io.BytesIO(); view.save(b,'JPEG',quality=92)
                    z.writestr(f'{i:03d}_{name}.jpg',b.getvalue())
                manifest.append({'caso':i,'id':k,'archivos':r['archivos'],'lectura':r.get('lectura'),'problemas':r.get('problemas',[])})
            z.writestr('lecturas.json',json.dumps(manifest,ensure_ascii=False,indent=2))
            z.writestr('INSTRUCCIONES.txt','Compara cada lectura con la página, encabezado y cargo. Verifica expediente, código y fecha manuscrita de notificación. No completes dígitos ilegibles. Devuelve una tabla con caso, expediente, código, fecha y observaciones. Las correcciones se confirman manualmente en el cuaderno antes de exportar.')
        print(dest)
        return dest

def _patch_xlsx(source,output,updates,epoch):
    """Modifica únicamente las celdas autorizadas y sus formatos. Conserva las otras partes ZIP."""
    parser=ET.XMLParser(resolve_entities=False,no_network=True)
    def parse(b): return ET.fromstring(b,parser)
    def serial(root): return ET.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    with zipfile.ZipFile(source) as zin:
        book=parse(zin.read('xl/workbook.xml'))
        rels=parse(zin.read('xl/_rels/workbook.xml.rels'))
        targets={e.get('Id'):posixpath.normpath(posixpath.join('xl',e.get('Target'))) if not e.get('Target').startswith('/') else e.get('Target').lstrip('/') for e in rels}
        sheet_paths={e.get('name'):targets[e.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')] for e in book.find('s:sheets',N)}
        style=parse(zin.read('xl/styles.xml'))
        xfs=style.find('s:cellXfs',N)
        formats=style.find('s:numFmts',N)
        if formats is None:
            formats=ET.Element('{'+NS+'}numFmts',count='0'); style.insert(0,formats)
        fmt_id=max([163]+[int(e.get('numFmtId')) for e in formats])+1
        ET.SubElement(formats,'{'+NS+'}numFmt',numFmtId=str(fmt_id),formatCode='dd/mm/yyyy')
        formats.set('count',str(len(formats)))
        style_cache={}
        def date_style(old):
            if old not in style_cache:
                xf=copy.deepcopy(xfs[old]); xf.set('numFmtId',str(fmt_id)); xf.set('applyNumberFormat','1')
                style_cache[old]=len(xfs); xfs.append(xf)
            return str(style_cache[old])
        changed={}
        grouped=defaultdict(dict)
        for (sheet,cell),d in updates.items(): grouped[sheet][cell]=d
        for sheet,cells in grouped.items():
            root=parse(zin.read(sheet_paths[sheet]))
            rows={int(r.get('r')):r for r in root.find('s:sheetData',N)}
            merged=root.find('s:mergeCells',N)
            if merged is not None:
                from openpyxl.utils.cell import range_boundaries
                for addr in cells:
                    col,row=re.fullmatch(r'([A-Z]+)(\d+)',addr).groups(); ci=column_index_from_string(col); ri=int(row)
                    for m in merged:
                        a,b,c,d=range_boundaries(m.get('ref'))
                        if a<=ci<=c and b<=ri<=d: raise ValueError('Celda de destino combinada: '+addr)
            for addr,d in cells.items():
                col,row_num=re.fullmatch(r'([A-Z]+)(\d+)',addr).groups()
                row=rows[int(row_num)]
                cell=next((c for c in row if c.get('r')==addr),None)
                if cell is None:
                    cell=ET.Element('{'+NS+'}c',r=addr)
                    ci=column_index_from_string(col)
                    pos=next((i for i,c in enumerate(row) if column_index_from_string(re.match('[A-Z]+',c.get('r')).group())>ci),len(row))
                    row.insert(pos,cell)
                if cell.find('s:f',N) is not None: raise ValueError('No se reemplaza una fórmula: '+addr)
                cell.set('s',date_style(int(cell.get('s','0')))); cell.set('t','n')
                for child in list(cell):
                    if child.tag in {'{'+NS+'}v','{'+NS+'}is'}: cell.remove(child)
                ET.SubElement(cell,'{'+NS+'}v').text=str(to_excel(datetime.combine(d,datetime.min.time()),epoch))
            changed[sheet_paths[sheet]]=serial(root)
        xfs.set('count',str(len(xfs)))
        changed['xl/styles.xml']=serial(style)
        tmp=Path(str(output)+'.tmp')
        try:
            with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as zout:
                for info in zin.infolist(): zout.writestr(info,changed.get(info.filename,zin.read(info.filename)))
            with zipfile.ZipFile(tmp) as check:
                bad=check.testzip()
                if bad: raise ValueError('Error ZIP: '+bad)
                for sheet,cells in grouped.items():
                    root=parse(check.read(sheet_paths[sheet]))
                    vals={c.get('r'):c.findtext('s:v',namespaces=N) for c in root.findall('s:sheetData/s:row/s:c',N) if c.get('r') in cells}
                    for addr,d in cells.items():
                        if float(vals[addr])!=to_excel(datetime.combine(d,datetime.min.time()),epoch):
                            raise ValueError('La verificación de exportación falló: '+addr)
            os.replace(tmp,output)
        finally:
            tmp.unlink(missing_ok=True)


def patch_xlsx(source, output, updates, epoch):
    # Segunda barrera: incluso una llamada directa debe respetar E, F y hoja.
    if Path(source).resolve() == Path(output).resolve():
        raise ValueError('La salida debe ser una copia distinta del original.')
    index, _, original_epoch = build_index(source, (HOJA,))
    allowed = {(r['hoja'],r['celda']):r for items in index.values() if len(items)==1
               for r in items if fecha(r['emision'])==FECHA_EMISION
               and r['fecha_actual'] in (None,'')}
    for key,d in updates.items():
        if key not in allowed:
            raise ValueError('Celda fuera del alcance o ya ocupada: '+str(key))
        if not isinstance(d,date) or not FECHA_EMISION<=d<=datetime.now(ZoneInfo('America/Lima')).date():
            raise ValueError('Fecha de notificación inválida: '+str(d))
    if not updates:
        raise ValueError('No hay fechas confirmadas para escribir.')
    _patch_xlsx(source,output,updates,original_epoch)
